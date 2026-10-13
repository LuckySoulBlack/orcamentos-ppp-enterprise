from pathlib import Path
import shutil
import unicodedata
from xml.sax.saxutils import escape
from datetime import datetime
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, A3, A2, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
try:
    from PIL import Image as PILImage
except Exception:
    PILImage = None

from core.database import conn, get_pasta_local_relatorios
from services.dashboard_html_service import gerar_dashboard_html

AZUL = '#0B3B75'
AZUL_ESCURO = '#071A33'
AZUL_CLARO = '#DBEAFE'
CINZA = '#F5F8FC'


def moeda(valor):
    try:
        valor = float(valor or 0)
    except Exception:
        valor = 0
    return f'R$ {valor:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def moeda4(valor):
    try:
        valor = float(valor or 0)
    except Exception:
        valor = 0
    return f'R$ {valor:,.4f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _limpar_nome_texto(v):
    s = str(v or '')
    s = s.replace('T#U00c9CNICA', 'TECNICA').replace('TÉCNICA', 'TECNICA').replace('TÃ‰CNICA', 'TECNICA')
    return s



def _resumir_texto(v, limite=140):
    # REGRA DOCUMENTAL PPP: documentos oficiais não podem cortar texto.
    # Mantém a assinatura da função para compatibilidade, mas retorna o texto integral.
    return _limpar_nome_texto(v or '').strip()


def _segmento_lower(orc):
    return _limpar_nome_texto(texto(orc, 'segmento_nome') or texto(orc, 'segmento') or '').lower()


def _municipio_uf(orc):
    municipio = texto(orc, 'cliente_municipio') or texto(orc, 'municipio') or texto(orc, 'empresa_municipio') or ''
    uf = texto(orc, 'cliente_uf') or texto(orc, 'uf') or texto(orc, 'empresa_uf') or ''
    if municipio and uf:
        return f"{municipio}/{uf}"
    return municipio or uf or ''


def _contrato_ref(orc):
    """Monta a referência contratual sem duplicar prefixos.

    Corrige casos em que o campo número já vem cadastrado como
    "Contrato nº 14/2023", evitando duplicidade de prefixo.
    """
    import re

    partes = []

    numero_contrato = _limpar_nome_texto(texto(orc, 'contrato_numero')).strip()
    nome_contrato = _limpar_nome_texto(texto(orc, 'contrato_nome')).strip()

    if numero_contrato:
        # Remove qualquer prefixo já existente no cadastro e aplica um único padrão.
        numero_limpo = re.sub(
            r'^(contrato\s+n[º°o]?\.?\s*|n[º°o]?\.?\s*)+',
            '',
            numero_contrato,
            flags=re.IGNORECASE
        ).strip(' -–—')
        if re.match(r'^\d', numero_limpo):
            partes.append(f"Contrato nº {numero_limpo}")
        else:
            partes.append(numero_contrato)

    if nome_contrato:
        # Se o nome do contrato começar repetindo o mesmo número, remove essa parte.
        if numero_contrato:
            numero_base = re.sub(
                r'^(contrato\s+n[º°o]?\.?\s*|n[º°o]?\.?\s*)+',
                '',
                numero_contrato,
                flags=re.IGNORECASE
            ).strip(' -–—')
            if numero_base:
                nome_contrato = re.sub(
                    rf'^(contrato\s*)?(n[º°o]?\.?\s*)?{re.escape(numero_base)}\s*[-–—]?\s*',
                    '',
                    nome_contrato,
                    flags=re.IGNORECASE
                ).strip()
        if nome_contrato:
            partes.append(nome_contrato)

    aditivo = texto(orc, 'aditivo_numero') or texto(orc, 'numero_aditivo')
    if aditivo:
        aditivo_txt = _limpar_nome_texto(aditivo).strip()
        if aditivo_txt.lower().startswith('aditivo'):
            partes.append(aditivo_txt)
        else:
            partes.append(f"Aditivo {aditivo_txt}")

    ref = ' - '.join([_limpar_nome_texto(p).strip() for p in partes if str(p).strip()])
    # Sanitização final contra qualquer duplicidade residual.
    ref = re.sub(r'Contrato\s+n[º°o]?\.?\s+Contrato\s+n[º°o]?\.?', 'Contrato nº', ref, flags=re.IGNORECASE)
    ref = re.sub(r'Contrato\s+n[º°o]?\.?\s+n[º°o]?\.?', 'Contrato nº', ref, flags=re.IGNORECASE)
    ref = re.sub(r'Contrato\s+n[º°o]?\.?\s+(Contrato\s+de\b)', r'\1', ref, flags=re.IGNORECASE)
    return ref or 'contrato/aditivo vinculado'


def _objeto_resumido(orc):
    objeto = texto(orc, 'objeto') or texto(orc, 'contrato_objeto') or ''
    return _resumir_texto(objeto, 180)


def _principais_etapas_capex(itens, limite=4):
    etapas = []
    vistos = set()
    for item in itens or []:
        if str(texto(item, 'capex_opex')).upper() != 'CAPEX':
            continue
        etapa = _limpar_nome_texto(texto(item, 'etapa') or texto(item, 'classificacao') or '')
        if etapa and etapa not in vistos:
            vistos.add(etapa)
            etapas.append(etapa)
        if len(etapas) >= limite:
            break
    if not etapas:
        return 'etapas e itens discriminados na planilha orçamentária'
    if len(etapas) == 1:
        return etapas[0]
    return ', '.join(etapas[:-1]) + ' e ' + etapas[-1]


def _tem_opex(itens=None, opex=None):
    if opex:
        try:
            return len(opex) > 0
        except Exception:
            pass
    for item in itens or []:
        if str(texto(item, 'capex_opex')).upper() != 'CAPEX':
            return True
    return False


def _tem_os_interna(anexos):
    for a in anexos or []:
        nome = str(texto(a, 'nome_arquivo') or texto(a, 'caminho') or '').lower()
        tipo = str(texto(a, 'tipo') or '').lower()
        obs = str(texto(a, 'observacao') or '').lower()
        if 'os' in nome or 'ordem de serviço' in tipo or 'ordem de servico' in tipo or 'os interna' in obs:
            return True
    return False


def _status_os_texto(anexos):
    bruto = ' '.join(str(texto(a, 'observacao') or '').lower() for a in (anexos or []))
    if 'finaliz' in bruto or 'conclu' in bruto or 'atesto' in bruto:
        return 'FINALIZADO'
    if 'execu' in bruto or 'andamento' in bruto:
        return 'EM EXECUÇÃO'
    if 'pendente' in bruto or 'a executar' in bruto:
        return 'PENDENTE'
    return ''


def gerar_descricao_financeira_capex(orc):
    objeto = _objeto_resumido(orc)
    if objeto:
        return f"CAPEX - {objeto}"

    seg = _segmento_lower(orc)
    if 'ilumina' in seg:
        return 'CAPEX - implantação/adequação de pontos de iluminação pública'
    if 'lpr' in seg or 'videomonitor' in seg or 'monitoramento' in seg:
        return 'CAPEX - implantação de sistema de monitoramento e leitura de placas'
    if 'solar' in seg or 'fotovolta' in seg or 'energia' in seg:
        return 'CAPEX - implantação de solução fotovoltaica'
    if 'telecom' in seg:
        return 'CAPEX - implantação/adequação de infraestrutura de telecomunicações'
    if 'cco' in seg:
        return 'CAPEX - implantação/adequação de centro de controle operacional'
    return 'CAPEX - implantação/adequação do objeto contratado'


def _fontes_resumo_financeiro(itens=None, opex=None, limite=4):
    fontes = []
    vistos = set()
    for row in list(itens or []) + list(opex or []):
        fonte = _limpar_nome_texto(texto(row, 'fonte') or '')
        if fonte and fonte.lower() not in vistos:
            vistos.add(fonte.lower())
            fontes.append(fonte)
        if len(fontes) >= limite:
            break
    return ', '.join(fontes) if fontes else 'fontes e composições cadastradas'


def criterio_financeiro_capex(orc, itens=None, anexos=None):
    etapas = _principais_etapas_capex(itens or [], limite=5)
    fontes = _fontes_resumo_financeiro(itens or [])
    criterio = (
        "Investimento CAPEX vinculado ao contrato/aditivo informado, composto por fornecimento, implantação, "
        "serviços técnicos, materiais, equipamentos, instalação, configuração, testes e comissionamento necessários "
        f"ao atendimento do objeto descrito na proposta. Critério formado pelas etapas CAPEX: {etapas}; "
        f"valores finais calculados conforme {fontes}."
    )
    if anexos:
        criterio += " Anexos, cotações, OS e registros técnicos atuam como evidências de suporte."
    return _resumir_texto(_limpar_nome_texto(criterio), 520)


def criterio_parcelamento_capex(qtd_parcelas, valor_parcela):
    if qtd_parcelas:
        return (
            f"Distribuição financeira do CAPEX em {qtd_parcelas} parcela(s) de {moeda(valor_parcela)}, "
            "condicionada ao cronograma físico-financeiro aprovado, execução efetiva, medição, aceite/atesto "
            "e regras contratuais ou administrativas aplicáveis."
        )
    return (
        "Parcelamento não detalhado no orçamento. A liberação financeira deve observar cronograma aprovado, "
        "medição dos quantitativos executados, aceite/atesto e regras contratuais ou administrativas aplicáveis."
    )


def criterio_financeiro_opex(itens=None, opex=None):
    fontes = _fontes_resumo_financeiro(itens or [], opex or [])
    tem = _tem_opex(itens, opex)
    if not tem:
        return "Não há OPEX mensal adicional previsto neste lançamento."
    categorias = []
    vistos = set()
    for row in opex or []:
        nome = _limpar_nome_texto(texto(row, 'categoria') or texto(row, 'centro_custo') or '')
        if nome and nome.lower() not in vistos:
            vistos.add(nome.lower())
            categorias.append(nome)
    for row in itens or []:
        if str(texto(row, 'capex_opex')).upper() == 'CAPEX':
            continue
        nome = _limpar_nome_texto(texto(row, 'classificacao') or texto(row, 'etapa') or '')
        if nome and nome.lower() not in vistos:
            vistos.add(nome.lower())
            categorias.append(nome)
    escopo = ', '.join(categorias[:5]) if categorias else 'operação, suporte, manutenção, monitoramento ou serviço continuado'
    return (
        f"Valor mensal recorrente previsto para {escopo}, calculado conforme {fontes}. "
        "Não compõe o parcelamento do CAPEX e deve ser tratado separadamente para fins de operação, medição mensal, atesto e faturamento recorrente."
    )


def gerar_texto_proposta_dinamico(orc, itens=None, opex=None, anexos=None):
    objeto = _objeto_resumido(orc) or 'o objeto descrito neste orçamento'
    municipio = _municipio_uf(orc)
    contrato = _contrato_ref(orc)
    segmento = texto(orc, 'segmento_nome') or texto(orc, 'segmento') or 'segmento vinculado'
    etapas = _principais_etapas_capex(itens or [])
    tem_opex = _tem_opex(itens, opex)
    tem_os = _tem_os_interna(anexos)
    status_os = _status_os_texto(anexos)

    local_txt = f", em {municipio}" if municipio else ""
    os_txt = ""
    if tem_os:
        if status_os == 'PENDENTE':
            os_txt = " O levantamento considera OS interna pendente, com informações previstas para análise, registros fotográficos, localização e relação de materiais, sem caracterizar serviço executado, medido ou aceito."
        elif status_os == 'EM EXECUÇÃO':
            os_txt = " O levantamento considera OS interna em execução, com registros técnicos, localização e relação de materiais para acompanhamento da intervenção."
        elif status_os == 'FINALIZADO':
            os_txt = " O levantamento considera OS interna finalizada e documentação técnica de suporte, quando acompanhada de registros fotográficos, atesto ou evidências de conclusão."
        else:
            os_txt = " O levantamento considera documentação técnica interna, com registros fotográficos, localização e relação de materiais quando anexados ao processo."

    seg_lower = _segmento_lower(orc)
    if 'ilumina' in seg_lower:
        finalidade = "O objeto caracteriza implantação, adequação ou melhoria do sistema de iluminação pública, visando segurança, visibilidade noturna, eficiência operacional e melhores condições de uso do espaço público."
        pos = "Após execução, validação técnica e aceite da fiscalização, os equipamentos e pontos implantados ou adequados poderão ser incorporados à rotina de operação da PPP, ficando a concessionária responsável pela manutenção corretiva conforme regras contratuais aplicáveis."
    elif 'lpr' in seg_lower or 'videomonitor' in seg_lower or 'monitoramento' in seg_lower:
        finalidade = "O objeto caracteriza implantação, adequação ou melhoria de infraestrutura de monitoramento, conectividade e apoio operacional, conforme escopo técnico e condições contratuais aplicáveis."
        pos = "Após execução, validação técnica e aceite da fiscalização, os equipamentos e sistemas implantados poderão ser incorporados à rotina operacional da PPP."
    elif 'solar' in seg_lower or 'fotovolta' in seg_lower or 'energia' in seg_lower:
        finalidade = "O objeto caracteriza implantação, adequação ou melhoria de solução de geração ou suporte energético, observadas as premissas técnicas do orçamento e as condições contratuais aplicáveis."
        pos = "Após execução, validação técnica e aceite da fiscalização, a solução implantada poderá ser incorporada à rotina operacional e de manutenção da PPP."
    else:
        finalidade = "O objeto caracteriza implantação, adequação ou melhoria de infraestrutura pública vinculada ao contrato, observadas as premissas técnicas, quantitativos e condições registradas no orçamento."
        pos = "Após execução, validação técnica e aceite da fiscalização, os serviços e equipamentos implantados poderão ser incorporados à rotina de operação e manutenção conforme regras contratuais aplicáveis."

    opex_txt = (
        "Este orçamento contempla CAPEX de implantação/adequação e OPEX mensal previsto, demonstrado separadamente, não compondo o cronograma físico-financeiro do CAPEX."
        if tem_opex else
        "Este orçamento contempla apenas CAPEX de implantação/adequação. Não há OPEX mensal adicional previsto neste lançamento, salvo posterior definição específica de remuneração, reequilíbrio ou alteração contratual."
    )

    return (
        f"O presente orçamento tem por finalidade registrar {objeto}{local_txt}, vinculado ao {contrato}, segmento {segmento}."
        f"{os_txt}\n\n"
        f"A intervenção contempla os materiais, equipamentos, mão de obra e serviços necessários, conforme etapas e itens discriminados na planilha orçamentária, incluindo {etapas}.\n\n"
        f"{finalidade}\n\n"
        f"{pos}\n\n"
        f"{opex_txt}"
    )


def gerar_texto_memorial_intro(orc, itens=None, opex=None):
    objeto = (_objeto_resumido(orc) or 'objeto cadastrado no orçamento').strip()
    # Evita dupla pontuação quando o objeto já vem finalizado com ponto no cadastro.
    objeto = objeto.rstrip(' .')
    etapas = _principais_etapas_capex(itens or [], limite=6)
    tem_opex = _tem_opex(itens, opex)
    opex_txt = ' Há também previsão de OPEX demonstrada em seção própria, sem composição do cronograma CAPEX.' if tem_opex else ' Não há OPEX mensal adicional previsto neste lançamento.'
    return (
        f"Este memorial descreve tecnicamente {objeto}, organizado a partir das etapas e itens lançados no orçamento. "
        f"A estrutura CAPEX contempla {etapas}, conforme quantidades, unidades, fontes de referência e condições registradas na planilha orçamentária."
        f"{opex_txt}"
    )


def _nome_arquivo_seguro(nome):
    s = _limpar_nome_texto(Path(str(nome or 'arquivo')).name)
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode('ascii')
    permitidos = []
    for ch in s:
        if ch.isalnum() or ch in (' ', '.', '_', '-'):
            permitidos.append(ch)
        else:
            permitidos.append('_')
    s = ''.join(permitidos).strip().replace(' ', '_')
    while '__' in s:
        s = s.replace('__', '_')
    return s or 'arquivo'


def _tipo_anexo_padrao(tipo, nome_arquivo=''):
    t = str(tipo or '').strip()
    nome = str(nome_arquivo or '').lower()

    if t and t.lower() not in ('outro', 'outros', 'documento'):
        return t

    if 'financeir' in nome or 'aporte' in nome:
        return 'Proposta financeira'
    if 'comercial' in nome:
        return 'Proposta comercial'
    if 'tecnica' in nome or 'técnica' in nome or 'fornecedor' in nome:
        return 'Proposta técnica de fornecedor'
    if 'contrato' in nome or 'concessao' in nome or 'concessão' in nome or 'anexo_iv' in nome or 'minuta' in nome:
        return 'Contrato / documento contratual'
    if 'encargo' in nome:
        return 'Caderno de Encargos'
    if 'ordem' in nome or 'os' in nome:
        return 'Ordem de Serviço'
    if 'aceite' in nome or 'atesto' in nome:
        return 'Aceite'
    if 'art' in nome or 'rrt' in nome:
        return 'ART/RRT'
    if 'nota' in nome or 'nf' in nome:
        return 'Nota fiscal'
    if 'foto' in nome or 'fotograf' in nome:
        return 'Relatório fotográfico'
    if 'medicao' in nome or 'medição' in nome:
        return 'Relatório de medição'
    return 'Documento complementar'


def validar_documentos_finais(orcamento_id):
    """Validações não destrutivas antes do pacote final."""
    avisos = []
    orc = buscar_orcamento(orcamento_id)
    itens = buscar_itens(orcamento_id)
    opex = buscar_opex(orcamento_id)
    cron = buscar_cronograma(orcamento_id)
    anexos = buscar_anexos(orcamento_id)

    capex = sum(total_item(i) for i in itens if str(texto(i, 'capex_opex')).upper() == 'CAPEX')
    total_cron = 0.0
    parcelas_pct = defaultdict(float)

    for row in cron:
        p = int(numero(texto(row, 'parcela')) or 0)
        pct = numero(texto(row, 'percentual'))
        if p > 0:
            parcelas_pct[p] += pct

    if capex <= 0:
        avisos.append('CAPEX final não localizado ou zerado.')
    if anexos is None:
        avisos.append('Falha ao consultar anexos.')
    if isinstance(anexos, list) and len(anexos) == 0:
        avisos.append('Nenhum anexo cadastrado para este orçamento.')

    # Checagem visual de itens: qtd x preço unitário final deve bater com total final.
    for item in itens:
        qtd = numero(texto(item, 'quantidade'))
        if qtd:
            unit = preco_unitario_final(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem'))
            total_calc = qtd * unit
            total_real = total_item(item)
            if abs(total_calc - total_real) > 0.01:
                avisos.append(f"Divergência visual no item {texto(item,'codigo')}: qtd x preço unitário final não bate com o total.")

    return avisos


def numero(valor):
    try:
        return float(valor or 0)
    except Exception:
        return 0.0


def texto(row, chave, padrao=''):
    try:
        return row[chave] if chave in row.keys() and row[chave] is not None else padrao
    except Exception:
        return padrao


def calcular_total_item(qtd, custo, bdi, margem):
    return numero(qtd) * numero(custo) * (1 + numero(bdi) / 100.0) * (1 + numero(margem) / 100.0)


def preco_unitario_final(qtd, custo, bdi, margem):
    total = calcular_total_item(qtd, custo, bdi, margem)
    q = numero(qtd)
    return total / q if q else 0.0


def total_item(item):
    return calcular_total_item(texto(item, 'quantidade'), texto(item, 'custo_unit'), texto(item, 'bdi'), texto(item, 'margem'))


def total_opex(row):
    return numero(texto(row, 'qtd_mensal')) * numero(texto(row, 'custo_unit')) * (1 + numero(texto(row, 'bdi')) / 100.0) * (1 + numero(texto(row, 'margem')) / 100.0)


def safe_codigo(codigo):
    codigo = str(codigo or 'ORCAMENTO').strip() or 'ORCAMENTO'
    for ch in ['/', '\\', ':', '*', '?', '"', '<', '>', '|']:
        codigo = codigo.replace(ch, '-')
    return codigo


def buscar_orcamento(orcamento_id):
    with conn() as c:
        return c.execute('''
            SELECT o.*, emp.razao AS empresa_razao, emp.fantasia AS empresa_fantasia, emp.cnpj AS empresa_cnpj,
            emp.endereco AS empresa_endereco, emp.municipio AS empresa_municipio, emp.uf AS empresa_uf,
            emp.representante AS empresa_representante, emp.cargo_representante AS empresa_cargo,
            emp.responsavel_tecnico AS empresa_responsavel_tecnico, emp.registro_tecnico AS empresa_registro_tecnico, emp.logo_path AS empresa_logo_path,
            cli.nome AS cliente_nome, cli.cnpj AS cliente_cnpj, cli.municipio AS cliente_municipio, cli.uf AS cliente_uf,
            cli.orgao_setor AS cliente_setor, cli.responsavel AS cliente_responsavel, cli.cargo AS cliente_cargo,
            con.nome AS contrato_nome, con.numero AS contrato_numero, con.objeto AS contrato_objeto,
            seg.nome AS segmento_nome
            FROM orcamentos o
            LEFT JOIN clientes cli ON cli.id = o.cliente_id
            LEFT JOIN contratos con ON con.id = o.contrato_id
            LEFT JOIN segmentos seg ON seg.id = o.segmento_id
            LEFT JOIN empresa emp ON emp.id = 1
            WHERE o.id=?
        ''', (orcamento_id,)).fetchone()


def buscar_itens(orcamento_id):
    with conn() as c:
        return c.execute('SELECT * FROM itens WHERE orcamento_id=?ORDER BY COALESCE(etapa_id,0), etapa, id', (orcamento_id,)).fetchall()


def buscar_opex(orcamento_id):
    with conn() as c:
        return c.execute('SELECT * FROM opex WHERE orcamento_id=?ORDER BY categoria, id', (orcamento_id,)).fetchall()


def buscar_cronograma(orcamento_id):
    with conn() as c:
        rows = []
        try:
            rows = c.execute('''
                SELECT cp.*, eo.nome AS etapa_nome
                FROM cronograma_parcelas cp
                LEFT JOIN etapas_orcamento eo ON eo.id = cp.etapa_id
                WHERE cp.orcamento_id=?
                ORDER BY COALESCE(eo.ordem,9999), COALESCE(cp.etapa_id,0), cp.parcela
            ''', (orcamento_id,)).fetchall()
            if rows:
                return rows
        except Exception:
            pass
        try:
            return c.execute('SELECT * FROM cronograma WHERE orcamento_id=?ORDER BY parcela, id', (orcamento_id,)).fetchall()
        except Exception:
            return []


def buscar_anexos(orcamento_id):
    with conn() as c:
        try:
            return c.execute('''
                SELECT id, tipo, nome_arquivo, caminho, data_upload, observacao
                FROM documentos
                WHERE orcamento_id=?
                ORDER BY id DESC
            ''', (orcamento_id,)).fetchall()
        except Exception:
            try:
                return c.execute(
                    'SELECT * FROM anexos WHERE orcamento_id=?ORDER BY data_upload DESC, id DESC',
                    (orcamento_id,)
                ).fetchall()
            except Exception:
                return []


def pasta_orcamento(codigo):
    base = get_pasta_local_relatorios()
    pasta = base / safe_codigo(codigo)
    pasta.mkdir(parents=True, exist_ok=True)
    return pasta


def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='TituloPPP', parent=styles['Title'], alignment=TA_CENTER, textColor=colors.HexColor(AZUL_ESCURO), fontSize=17, leading=21, spaceAfter=10))
    styles.add(ParagraphStyle(name='SubtituloPPP', parent=styles['Normal'], alignment=TA_CENTER, textColor=colors.HexColor('#334155'), fontSize=8.5, leading=11, spaceAfter=8))
    styles.add(ParagraphStyle(name='SecaoPPP', parent=styles['Heading2'], textColor=colors.HexColor(AZUL), fontSize=12, leading=15, spaceBefore=8, spaceAfter=6))
    styles.add(ParagraphStyle(name='TextoPPP', parent=styles['Normal'], fontSize=8.8, leading=12, spaceAfter=5, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name='TextoTabela', parent=styles['Normal'], fontSize=6.2, leading=7.4, wordWrap='CJK', splitLongWords=True))
    styles.add(ParagraphStyle(name='TextoPequeno', parent=styles['Normal'], fontSize=7.2, leading=9.0, spaceAfter=3, wordWrap='CJK', splitLongWords=True))
    return styles


def _logo_path(orc):
    candidatos = []
    for chave in ('empresa_logo_path', 'logo_path'):
        try:
            val = texto(orc, chave)
            if val:
                candidatos.append(val)
        except Exception:
            pass
    try:
        with conn() as c:
            row = c.execute("SELECT logo_path FROM empresa WHERE id=1").fetchone()
            if row and row['logo_path']:
                candidatos.append(row['logo_path'])
    except Exception:
        pass
    try:
        with conn() as c:
            for row in c.execute("SELECT valor FROM config WHERE chave IN ('logo_path','empresa_logo_path','logo_empresa')"):
                if row['valor']:
                    candidatos.append(row['valor'])
    except Exception:
        pass

    for caminho in candidatos:
        try:
            p = Path(str(caminho).strip().strip('"'))
            if p.exists() and p.is_file():
                return str(p)
        except Exception:
            pass
    return None


def _imagem_logo(path, max_w=110, max_h=48):
    try:
        if PILImage is not None:
            with PILImage.open(path) as im:
                w, h = im.size
            if w <= 0 or h <= 0:
                return None
            escala = min(max_w / w, max_h / h)
            img = Image(path, width=w * escala, height=h * escala)
        else:
            img = Image(path, width=max_w, height=max_h)
        img.hAlign = 'CENTER'
        return img
    except Exception:
        return None


def _cabecalho_documento(story, orc, titulo, styles):
    logo = _logo_path(orc)
    if logo:
        img = _imagem_logo(logo)
        if img:
            story += [img, Spacer(1, 16)]
    story.append(Paragraph(titulo, styles['TituloPPP']))
    empresa = texto(orc, 'empresa_razao') or texto(orc, 'empresa_fantasia')
    cliente = texto(orc, 'cliente_nome')
    linha = f"{empresa or 'Empresa'} • {cliente or 'Cliente'} • {texto(orc,'cliente_municipio') or texto(orc,'municipio')}/{texto(orc,'cliente_uf') or texto(orc,'empresa_uf')}"
    story.append(Paragraph(linha, styles['SubtituloPPP']))
    story.append(Spacer(1, 5))


def _table_style(header=AZUL, font_size=7, total_bg=AZUL_CLARO):
    return TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(header)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#9AA7B8')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), font_size),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor(CINZA)]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(total_bg)),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('ALIGN', (5, 1), (-1, -1), 'RIGHT'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ])


def _resumos(itens, opex):
    capex = sum(total_item(i) for i in itens if str(texto(i, 'capex_opex')).upper() == 'CAPEX')
    opex_itens = sum(total_item(i) for i in itens if str(texto(i, 'capex_opex')).upper() != 'CAPEX')
    opex_rec = sum(total_opex(i) for i in opex)
    return capex, opex_itens + opex_rec, capex + opex_itens + opex_rec


def gerar_excel_revisao(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id); anexos = buscar_anexos(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '00_Revisao_Interna'; pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / 'Revisao_Financeira_Enterprise.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Resumo Executivo'
    azul = PatternFill('solid', fgColor='0B3B75'); azul_claro = PatternFill('solid', fgColor='DBEAFE'); cinza = PatternFill('solid', fgColor='F5F8FC')
    branco = Font(color='FFFFFF', bold=True); negrito = Font(bold=True, color='071A33')
    borda = Border(left=Side(style='thin', color='D9E2F3'), right=Side(style='thin', color='D9E2F3'), top=Side(style='thin', color='D9E2F3'), bottom=Side(style='thin', color='D9E2F3'))
    capex, opex_total, total = _resumos(itens, opex)
    ws.merge_cells('A1:D1'); ws['A1'] = 'REVISÃO INTERNA DO ORÇAMENTO PPP'; ws['A1'].font = Font(color='FFFFFF', bold=True, size=14); ws['A1'].fill = azul; ws['A1'].alignment = Alignment(horizontal='center')
    linhas = [['Código', texto(orc,'codigo'), 'Status', texto(orc,'status')], ['Cliente', texto(orc,'cliente_nome'), 'UF', texto(orc,'cliente_uf')], ['Objeto', texto(orc,'objeto'), 'Segmento', texto(orc,'segmento_nome')], ['CAPEX', capex, 'OPEX', opex_total], ['TOTAL GERAL', total, 'Data', datetime.now().strftime('%d/%m/%Y %H:%M')]]
    for i, linha in enumerate(linhas, start=3):
        for j, val in enumerate(linha, start=1):
            ws.cell(i,j,val); ws.cell(i,j).border=borda; ws.cell(i,j).alignment=Alignment(wrap_text=True, vertical='top')
            if j in (1,3): ws.cell(i,j).font=negrito; ws.cell(i,j).fill=azul_claro
            if isinstance(val,(int,float)): ws.cell(i,j).number_format='R$ #,##0.00'
    ws2 = wb.create_sheet('Planilha Orçamentária')
    headers = ['Etapa','Fonte','Código','Descrição','UN','Qtd','Custo Unit.','BDI %','Preço Total','Tipo','Classificação']
    ws2.append(headers)
    for cell in ws2[1]: cell.font=branco; cell.fill=azul; cell.alignment=Alignment(horizontal='center')
    for item in itens:
        ws2.append([texto(item,'etapa'), texto(item,'fonte'), texto(item,'codigo'), texto(item,'descricao'), texto(item,'unidade'), numero(texto(item,'quantidade')), numero(texto(item,'custo_unit')), numero(texto(item,'bdi')), total_item(item), texto(item,'capex_opex'), texto(item,'classificacao')])
    for row in opex:
        ws2.append(['OPEX '+str(texto(row,'categoria')), texto(row,'fonte'), texto(row,'codigo'), texto(row,'descricao'), texto(row,'unidade'), numero(texto(row,'qtd_mensal')), numero(texto(row,'custo_unit')), numero(texto(row,'bdi')), total_opex(row), 'OPEX', texto(row,'centro_custo')])
    ws2.append(['TOTAL','','','','','','','',total,'',''])
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.border=borda; cell.alignment=Alignment(wrap_text=True, vertical='top')
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18 if col != 4 else 52
    wb.save(arquivo); return arquivo


def _quebrar_texto_longo(v):
    s = str(v or '')
    s = s.replace('\\', '/')
    s = s.replace('/', '/<br/>')
    s = s.replace('_', '_<br/>')
    s = s.replace('-', '-<br/>')
    return s


def _paragraph(v, styles):
    # Paragraph com escape para preservar texto integral e permitir quebra automática.
    txt = _limpar_nome_texto(v or '')
    txt = escape(str(txt)).replace('\n', '<br/>')
    return Paragraph(txt, styles['TextoTabela'])


def _paragraph_bold(v, styles):
    txt = escape(_limpar_nome_texto(v or '')).replace('\n', '<br/>')
    return Paragraph(f'<b>{txt}</b>', styles['TextoTabela'])


def _obs_item(row):
    return _limpar_nome_texto(
        texto(row, 'observacao')
        or texto(row, 'observacoes')
        or texto(row, 'obs')
        or texto(row, 'justificativa')
        or ''
    ).strip()


def gerar_pdf_planilha_final(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '01_Cliente_Final'; pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / '02_Planilha_Orcamentaria_PPP.pdf'; styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=landscape(A4), rightMargin=10, leftMargin=10, topMargin=14, bottomMargin=14)
    story = []
    _cabecalho_documento(story, orc, 'PLANILHA ORÇAMENTÁRIA PPP', styles)
    story += [
        Paragraph(f"Objeto: {_limpar_nome_texto(texto(orc,'objeto'))}", styles['TextoPPP']),
        Paragraph(
            f"Referência UF: {texto(orc,'cliente_uf') or texto(orc,'empresa_uf')} • Fonte: SINAPI / composição própria / cotação, conforme cadastro do item. "
            "Os valores unitários apresentados correspondem ao preço final aplicado no orçamento.",
            styles['TextoPPP']
        ),
        Spacer(1, 6)
    ]

    def qfmt(v):
        return f"{numero(v):,.4f}".replace(',', 'X').replace('.', ',').replace('X','.')

    def montar_tabela_capex():
        capex_rows = [i for i in itens if str(texto(i, 'capex_opex')).upper() == 'CAPEX']
        data = [['Item','Etapa / Classificação','Fonte','Código','Descrição','UN','Qtd','Preço Unit. Final','Preço Total']]
        item_idx = 1; total_capex = 0.0
        grupos = OrderedDict()
        for item in capex_rows:
            grupos.setdefault(texto(item, 'etapa') or 'Etapa não informada', []).append(item)
        for etapa, rows in grupos.items():
            subtotal = sum(total_item(i) for i in rows); total_capex += subtotal
            data.append([
                '', _paragraph_bold(etapa, styles), '', '',
                _paragraph_bold('Subtotal CAPEX da etapa', styles),
                '', '', '', moeda(subtotal)
            ])
            for item in rows:
                unit_final = preco_unitario_final(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem'))
                data.append([
                    item_idx,
                    _paragraph(_limpar_nome_texto(texto(item,'classificacao') or 'CAPEX'), styles),
                    _paragraph(_limpar_nome_texto(texto(item,'fonte')), styles),
                    _paragraph(_limpar_nome_texto(texto(item,'codigo')), styles),
                    _paragraph(_limpar_nome_texto(texto(item,'descricao')), styles),
                    texto(item,'unidade'),
                    qfmt(texto(item,'quantidade')),
                    moeda4(unit_final),
                    moeda(total_item(item))
                ])
                item_idx += 1
        if len(data) == 1:
            data.append(['-','-','-','-','Nenhum item CAPEX lançado.','-','-','-','-'])
        data.append(['TOTAL CAPEX','','','','','','','',moeda(total_capex)])
        table = Table(data, repeatRows=1, colWidths=[24,78,64,74,332,30,48,70,72])
        table.setStyle(_table_style(font_size=5.6))
        return table, total_capex

    def montar_tabela_opex():
        data = [['Item','Categoria / Centro','Fonte','Código','Descrição','UN','Qtd Mensal','Preço Unit. Final','Total Mensal']]
        total_opex_val = 0.0; idx = 1
        opex_itens = [i for i in itens if str(texto(i, 'capex_opex')).upper() != 'CAPEX']
        for item in opex_itens:
            valor_total = total_item(item); total_opex_val += valor_total
            unit_final = preco_unitario_final(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem'))
            data.append([
                idx, _paragraph(_limpar_nome_texto(texto(item,'classificacao') or texto(item,'etapa') or 'OPEX'), styles),
                _paragraph(_limpar_nome_texto(texto(item,'fonte')), styles), _paragraph(_limpar_nome_texto(texto(item,'codigo')), styles),
                _paragraph(_limpar_nome_texto(texto(item,'descricao')), styles), texto(item,'unidade'),
                qfmt(texto(item,'quantidade')), moeda4(unit_final), moeda(valor_total)
            ])
            idx += 1
        for row in opex:
            valor_total = globals()['total_opex'](row); total_opex_val += valor_total
            unit_final = preco_unitario_final(texto(row,'qtd_mensal'), texto(row,'custo_unit'), texto(row,'bdi'), texto(row,'margem'))
            data.append([
                idx, _paragraph(_limpar_nome_texto(texto(row,'categoria') or texto(row,'centro_custo') or 'OPEX'), styles),
                _paragraph(_limpar_nome_texto(texto(row,'fonte')), styles), _paragraph(_limpar_nome_texto(texto(row,'codigo')), styles),
                _paragraph(_limpar_nome_texto(texto(row,'descricao')), styles), texto(row,'unidade'),
                qfmt(texto(row,'qtd_mensal')), moeda4(unit_final), moeda(valor_total)
            ])
            idx += 1
        if len(data) == 1:
            data.append(['-','-','-','-','Nenhum item OPEX lançado.','-','-','-','-'])
        data.append(['TOTAL OPEX MENSAL','','','','','','','',moeda(total_opex_val)])
        table = Table(data, repeatRows=1, colWidths=[24,78,64,74,332,30,54,70,72])
        table.setStyle(_table_style(font_size=5.6))
        return table, total_opex_val

    capex_table, total_capex = montar_tabela_capex()
    story += [Paragraph('1. CAPEX - INVESTIMENTO / IMPLANTAÇÃO', styles['SecaoPPP']), capex_table, Spacer(1, 10)]
    opex_table, total_opex_valor = montar_tabela_opex()
    story += [Paragraph('2. OPEX - OPERAÇÃO / RECORRÊNCIA', styles['SecaoPPP']), opex_table, Spacer(1, 8)]
    resumo = Table([
        ['Resumo financeiro do orçamento', 'Valor'],
        ['CAPEX - implantação/adequação do objeto contratado', moeda(total_capex)],
        ['OPEX mensal previsto', moeda(total_opex_valor) + '/mês'],
    ], colWidths=[420,120])
    resumo.setStyle(_table_style(font_size=8))
    story.append(resumo)
    doc.build(story); return arquivo


def _numerar_paginas_proposta(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor('#64748B'))
    canvas.drawRightString(A4[0] - 36, 18, f"Página {doc.page}")
    canvas.restoreState()


def gerar_pdf_proposta_final(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '01_Cliente_Final'; pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / '01_Proposta_Tecnica_Comercial_PPP.pdf'; styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=30, bottomMargin=30)
    capex, opex_total, total = _resumos(itens, opex)

    cron = buscar_cronograma(orcamento_id)
    parcelas_ativas = sorted({
        int(numero(texto(r, 'parcela')) or 0)
        for r in cron
        if int(numero(texto(r, 'parcela')) or 0) > 0 and (numero(texto(r, 'percentual')) > 0 or numero(texto(r, 'valor')) > 0)
    })
    qtd_parcelas = len(parcelas_ativas)
    valor_parcela = capex / qtd_parcelas if qtd_parcelas else 0
    parcelamento_txt = f"{qtd_parcelas} parcelas de {moeda(valor_parcela)}" if qtd_parcelas else "Conforme cronograma físico-financeiro"

    story = []
    _cabecalho_documento(story, orc, 'PROPOSTA TÉCNICA E ECONÔMICA - PPP', styles)
    story += [
        Paragraph('1. Identificação da proposta', styles['SecaoPPP']),
        Paragraph(f"<b>Código:</b> {texto(orc,'codigo')}", styles['TextoPPP']),
        Paragraph(f"<b>Cliente:</b> {texto(orc,'cliente_nome')}", styles['TextoPPP']),
        Paragraph(f"<b>A/C:</b> {texto(orc,'cliente_responsavel') or '-'}", styles['TextoPPP']),
        Paragraph(f"<b>Setor:</b> {texto(orc,'cliente_setor') or '-'}", styles['TextoPPP']),
        Paragraph(f"<b>Município/UF:</b> {texto(orc,'cliente_municipio') or texto(orc,'municipio')}/{texto(orc,'cliente_uf') or texto(orc,'empresa_uf')}", styles['TextoPPP']),
        Paragraph(f"<b>Contrato:</b> {texto(orc,'contrato_numero')} - {texto(orc,'contrato_nome')}", styles['TextoPPP']),
        Paragraph(f"<b>Segmento:</b> {texto(orc,'segmento_nome')}", styles['TextoPPP']),
        Paragraph(f"<b>Validade:</b> {texto(orc,'validade')}", styles['TextoPPP']),
        Spacer(1, 12),
        Paragraph('2. Objeto', styles['SecaoPPP']),
        Paragraph(_limpar_nome_texto(texto(orc,'objeto') or texto(orc,'contrato_objeto') or 'Objeto conforme cadastro do orçamento.'), styles['TextoPPP']),
        Paragraph('3. Escopo técnico resumido', styles['SecaoPPP']),
        Paragraph('A proposta contempla fornecimento, mão de obra, serviços especializados, instalação, testes, comissionamento, registros de execução e demais atividades necessárias à entrega do objeto, observadas as condições cadastradas no orçamento e os documentos anexos ao processo.', styles['TextoPPP']),
    ]
    anexos = buscar_anexos(orcamento_id)
    texto_proposta_final = _limpar_nome_texto(texto(orc, 'texto_proposta')) or gerar_texto_proposta_dinamico(orc, itens, opex, anexos)
    story += [
        Paragraph('4. Considerações técnicas da proposta', styles['SecaoPPP']),
        Paragraph(texto_proposta_final.replace('\n', '<br/>'), styles['TextoPPP'])
    ]

    story += [Paragraph('5. Referência financeira do orçamento', styles['SecaoPPP'])]
    criterio_capex = criterio_financeiro_capex(orc, itens, anexos)
    criterio_parcela = criterio_parcelamento_capex(qtd_parcelas, valor_parcela)
    criterio_opex = criterio_financeiro_opex(itens, opex)
    resumo = Table([
        ['Descrição', 'Critério', 'Valor'],
        ['CAPEX', Paragraph(criterio_capex, styles['TextoPequeno']), moeda(capex)],
        ['Parcelamento CAPEX', Paragraph(criterio_parcela, styles['TextoPequeno']), moeda(capex)],
        ['OPEX mensal previsto', Paragraph(criterio_opex, styles['TextoPequeno']), moeda(opex_total) + '/mês'],
    ], colWidths=[115,270,75])
    resumo.setStyle(_table_style(font_size=8))

    assinaturas = Table([
        ['Responsável Técnico', 'Representante Legal'],
        ['____________________________________', '____________________________________'],
        [texto(orc,'empresa_responsavel_tecnico') or 'Responsável Técnico', texto(orc,'empresa_representante') or 'Representante legal'],
        [texto(orc,'empresa_registro_tecnico') or '', texto(orc,'empresa_razao') or texto(orc,'empresa_fantasia') or '']
    ], colWidths=[230,230])
    assinaturas.setStyle(TableStyle([
        ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),8),
        ('TOPPADDING',(0,0),(-1,-1),5),
        ('BOTTOMPADDING',(0,0),(-1,-1),5),
    ]))

    story += [
        resumo,
        Spacer(1,14),
        Paragraph('6. Premissas técnicas e orçamentárias', styles['SecaoPPP']),
        Paragraph('Os preços apresentados correspondem aos valores finais aplicáveis ao orçamento, conforme composição cadastrada e condições comerciais vinculadas à proposta.', styles['TextoPPP']),
        Paragraph('7. Condições de execução e medição', styles['SecaoPPP']),
        Paragraph('A execução e a medição deverão observar o cronograma físico-financeiro aprovado, os quantitativos efetivamente executados, documentos de aceite, registros fotográficos, ordens de serviço, notas fiscais, cotações, memórias de cálculo e demais evidências anexadas ao processo.', styles['TextoPPP']),
        Spacer(1, 12),
        KeepTogether([
            Paragraph('8. Assinaturas', styles['SecaoPPP']),
            Spacer(1, 10),
            assinaturas
        ])
    ]
    doc.build(story, onFirstPage=_numerar_paginas_proposta, onLaterPages=_numerar_paginas_proposta); return arquivo



def _html(v):
    return escape(_limpar_nome_texto(v or '')).replace('\n', '<br/>')


def _p(v, styles, style_name='TextoPPP'):
    return Paragraph(_html(v), styles[style_name])


def _p_bold(v, styles, style_name='TextoPPP'):
    return Paragraph(f"<b>{_html(v)}</b>", styles[style_name])


def _fmt_qtd_memorial(v):
    n = numero(v)
    if abs(n - int(n)) < 0.000001:
        return str(int(n))
    return f"{n:,.4f}".replace(',', 'X').replace('.', ',').replace('X', '.').rstrip('0').rstrip(',')


def _is_sinapi_row(row):
    fonte = (_limpar_nome_texto(texto(row, 'fonte')) or '').lower()
    codigo = (_limpar_nome_texto(texto(row, 'codigo')) or '').lower()
    classificacao = (_limpar_nome_texto(texto(row, 'classificacao')) or '').lower()
    return 'sinapi' in fonte or 'sinapi' in codigo or 'sinapi' in classificacao


def _criterio_por_fonte(fonte):
    f = (_limpar_nome_texto(fonte) or '').lower()
    if not f:
        return 'Fonte não informada no cadastro do item'
    if 'sinapi' in f:
        return 'Base oficial cadastrada no orçamento'
    if 'sicro' in f:
        return 'Base oficial cadastrada no orçamento'
    if 'compos' in f:
        return 'Composição própria cadastrada no orçamento'
    if 'cota' in f:
        return 'Cotação ou referência de mercado cadastrada'
    if 'fornec' in f:
        return 'Referência de fornecedor cadastrada'
    if 'contrat' in f:
        return 'Referência contratual cadastrada'
    if 'intern' in f or 'banco' in f:
        return 'Banco interno ou referência interna cadastrada'
    if 'estudo' in f:
        return 'Estudo técnico-comercial cadastrado'
    return 'Fonte cadastrada no item'


def _status_evidencia_linhas(rows):
    rows = list(rows or [])
    if not rows:
        return 'Sem itens vinculados'
    if any(_obs_item(r) for r in rows):
        return 'Observação técnica cadastrada'
    if all(_is_sinapi_row(r) for r in rows):
        return 'Fonte oficial cadastrada'
    return 'Sem justificativa complementar cadastrada'


def _table_memorial(data, col_widths, styles, font_size=7.2):
    data_paragraph = []
    for row in data:
        linha = []
        for cell in row:
            if hasattr(cell, 'wrap'):
                linha.append(cell)
            else:
                linha.append(Paragraph(_html(cell), styles['TextoPequeno']))
        data_paragraph.append(linha)
    table = Table(data_paragraph, repeatRows=1, colWidths=col_widths, splitByRow=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(AZUL)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), font_size),
        ('GRID', (0, 0), (-1, -1), 0.30, colors.HexColor('#9AA7B8')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(CINZA)]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    return table


def _add_secao(story, n, titulo, styles):
    story.append(Paragraph(f'{n}. {titulo}', styles['SecaoPPP']))


def _add_item_memorial(story, item, styles):
    """Adiciona item do memorial em bloco visual leve e flexível.

    Mantém o texto integral, sem cortes, mas organiza título, metadados e
    observação em uma tabela-card para melhorar a leitura em análise oficial.
    """
    desc = _limpar_nome_texto(texto(item, 'descricao') or 'Item cadastrado sem descrição técnica.')
    meta = []
    if texto(item, 'fonte'):
        meta.append(f"Fonte: {texto(item,'fonte')}")
    if texto(item, 'codigo'):
        meta.append(f"Código: {texto(item,'codigo')}")
    if texto(item, 'unidade'):
        meta.append(f"Unidade: {texto(item,'unidade')}")
    if texto(item, 'quantidade'):
        meta.append(f"Quantidade: {_fmt_qtd_memorial(texto(item,'quantidade'))}")
    if texto(item, 'classificacao'):
        meta.append(f"Classificação: {texto(item,'classificacao')}")

    linhas = [
        [Paragraph(f"<b>{_html(desc)}</b>", styles['TextoPPP'])],
    ]
    if meta:
        linhas.append([Paragraph(f"<font size='7'>{_html(' • '.join(meta))}</font>", styles['TextoPPP'])])

    obs = _obs_item(item)
    if obs:
        linhas.append([Paragraph(f"<font size='7'><b>Observação técnica / justificativa do item:</b><br/>{_html(obs)}</font>", styles['TextoPPP'])])
    elif not _is_sinapi_row(item):
        linhas.append([Paragraph(
            "<font size='7'><b>Observação técnica / justificativa do item:</b><br/>Item sem observação específica cadastrada. Recomenda-se complementar as premissas técnicas antes do protocolo oficial, quando aplicável.</font>",
            styles['TextoPPP']
        )])

    card = Table(linhas, colWidths=[490], splitByRow=1)
    card.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#FFFFFF')),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LINEBELOW', (0, 0), (-1, 0), 0.20, colors.HexColor('#E2E8F0')),
    ]))
    story.append(card)
    story.append(Spacer(1, 7))


def _add_observacoes_gerais_estruturadas(story, sec, obs_gerais, styles):
    """Estrutura observações gerais em blocos, sem alterar a redação cadastrada.

    Quando houver vários parágrafos, distribui nos subtítulos institucionais.
    Se houver apenas um parágrafo, mantém um único bloco para não inventar
    classificação textual.
    """
    texto_obs = _limpar_nome_texto(obs_gerais or '').strip()
    if not texto_obs:
        return
    _add_secao(story, sec, 'Observações técnicas gerais', styles)
    # Divide somente por parágrafos já existentes, sem reescrever o texto cadastrado.
    partes = [p.strip() for p in texto_obs.replace('\r\n', '\n').split('\n') if p.strip()]
    if len(partes) <= 1:
        story.append(_p(texto_obs, styles))
        return
    subtitulos = ['Escopo geral', 'Premissas adotadas', 'Compatibilização técnica', 'Condições complementares']
    for idx, parte in enumerate(partes, start=1):
        titulo = subtitulos[idx - 1] if idx <= len(subtitulos) else f'Complemento {idx}'
        story.append(Paragraph(f'<b>{sec}.{idx} {titulo}</b>', styles['TextoPPP']))
        story.append(_p(parte, styles))
        story.append(Spacer(1, 3))


def gerar_pdf_memorial_descritivo(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id); anexos = buscar_anexos(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '01_Cliente_Final'; pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / '03_Memorial_Descritivo_PPP.pdf'; styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=A4, rightMargin=38, leftMargin=38, topMargin=34, bottomMargin=34)

    etapas = OrderedDict()
    for item in itens:
        etapas.setdefault(texto(item,'etapa') or 'Etapa não informada', []).append(item)

    capex, opex_total, total_geral = _resumos(itens, opex)
    tem_opex = _tem_opex(itens, opex)
    objeto = _limpar_nome_texto(texto(orc,'objeto') or texto(orc,'contrato_objeto') or 'Objeto conforme cadastro do orçamento.')
    cliente = texto(orc, 'cliente_nome') or 'Cliente não informado'
    municipio_uf = _municipio_uf(orc) or 'Município/UF não informado'
    segmento = texto(orc, 'segmento_nome') or texto(orc, 'segmento') or 'Segmento não informado'
    contrato = _contrato_ref(orc)
    anexos = anexos or []

    story = []
    _cabecalho_documento(story, orc, 'MEMORIAL DESCRITIVO - PPP', styles)

    sec = 1
    _add_secao(story, sec, 'Finalidade', styles); sec += 1
    story.append(Paragraph(
        'Este memorial tem por finalidade complementar a planilha orçamentária, apresentando o escopo técnico, as fontes cadastradas, as premissas lançadas nos itens, a rastreabilidade documental e os critérios de medição relacionados ao orçamento. O documento foi estruturado para apoiar a análise técnica, administrativa e de controle, sem substituir projeto executivo, ordem de serviço, laudo, autorização formal, relatório de medição ou atesto de execução quando exigíveis.',
        styles['TextoPPP']
    ))

    _add_secao(story, sec, 'Objeto', styles); sec += 1
    story.append(_p(objeto, styles))

    _add_secao(story, sec, 'Identificação e base documental do orçamento', styles); sec += 1
    base_data = [
        ['Campo', 'Informação'],
        ['Código do orçamento', texto(orc,'codigo') or 'Não informado'],
        ['Contrato / referência', contrato],
        ['Cliente / órgão', cliente],
        ['Município/UF', municipio_uf],
        ['Segmento', segmento],
        ['CAPEX previsto', moeda(capex)],
        ['OPEX mensal previsto', moeda(opex_total) + '/mês' if tem_opex else 'Não previsto'],
        ['Documentos anexos identificados', str(len(anexos))],
    ]
    story.append(_table_memorial(base_data, [145, 345], styles, font_size=7.4))
    story.append(Spacer(1, 6))

    _add_secao(story, sec, 'Premissas metodológicas', styles); sec += 1
    premissas = [
        ['Premissa', 'Aplicação no memorial'],
        ['Origem dos dados', 'As informações são extraídas do cadastro do orçamento, das etapas, dos itens, das observações técnicas e dos anexos vinculados ao processo.'],
        ['Justificativa técnica', 'A justificativa técnica deve ser registrada no campo Observação do item. O memorial reproduz essa informação sem criar justificativas automáticas.'],
        ['Fontes de preço', 'A fonte considerada é aquela efetivamente cadastrada no item, podendo abranger base oficial, composição própria, cotação, fornecedor, contrato, banco interno, estudo técnico-comercial, referência de mercado ou outra fonte documental aplicável.'],
        ['Itens SINAPI', 'Quando o item estiver identificado como SINAPI, o memorial registra a fonte e o código cadastrados. A ausência de observação não gera alerta automático.'],
        ['Itens não SINAPI', 'Quando o item não estiver identificado como SINAPI, a observação técnica é recomendada para demonstrar premissas, composição, fornecedor, cotação ou justificativa do item.'],
    ]
    story.append(_table_memorial(premissas, [130, 360], styles, font_size=7.1))
    story.append(Spacer(1, 6))

    _add_secao(story, sec, 'Critério de preços e rastreabilidade', styles); sec += 1
    story.append(Paragraph(
        'A rastreabilidade do orçamento considera a fonte de referência efetivamente cadastrada em cada item. A justificativa técnica não é presumida pelo sistema: ela deve constar no campo Observação do item ou em documentação complementar anexada. Itens sem fonte ou sem justificativa, quando não vinculados a base oficial, devem ser complementados antes do protocolo oficial, conforme a natureza do objeto e o nível de exigência do processo.',
        styles['TextoPPP']
    ))
    rastreabilidade = [['Etapa', 'Fonte cadastrada', 'Itens', 'Critério de referência', 'Evidência / justificativa']]
    for etapa, rows in etapas.items():
        fontes = []
        for r in rows:
            f = _limpar_nome_texto(texto(r, 'fonte') or 'Fonte não informada')
            if f not in fontes:
                fontes.append(f)
        fonte_txt = ', '.join(fontes) if fontes else 'Fonte não informada'
        criterios = []
        for f in fontes:
            c = _criterio_por_fonte(f)
            if c not in criterios:
                criterios.append(c)
        rastreabilidade.append([
            etapa,
            fonte_txt,
            str(len(rows)),
            '; '.join(criterios) if criterios else 'Fonte não informada no cadastro do item',
            _status_evidencia_linhas(rows)
        ])
    if len(rastreabilidade) == 1:
        rastreabilidade.append(['Nenhuma etapa lançada', '-', '0', '-', '-'])
    story.append(_table_memorial(rastreabilidade, [125, 90, 35, 125, 115], styles, font_size=6.6))
    story.append(Spacer(1, 6))

    _add_secao(story, sec, 'Escopo consolidado por etapa', styles); sec += 1
    escopo = [['Etapa', 'Qtde. itens', 'Fontes identificadas', 'Valor CAPEX/OPEX calculado']]
    for etapa, rows in etapas.items():
        fontes = []
        for r in rows:
            f = _limpar_nome_texto(texto(r, 'fonte') or 'Fonte não informada')
            if f not in fontes:
                fontes.append(f)
        escopo.append([
            etapa,
            str(len(rows)),
            ', '.join(fontes) if fontes else 'Fonte não informada',
            moeda(sum(total_item(r) for r in rows))
        ])
    if len(escopo) == 1:
        escopo.append(['Nenhuma etapa lançada', '0', '-', moeda(0)])
    story.append(_table_memorial(escopo, [185, 55, 150, 100], styles, font_size=7.0))
    story.append(Spacer(1, 6))

    _add_secao(story, sec, 'Etapas, itens e justificativas técnicas', styles)
    sec_etapas = sec; sec += 1
    story.append(Paragraph(
        'A seguir são apresentados os itens do orçamento com descrição, fonte cadastrada, código, unidade, quantidade e, quando informada, a observação técnica cadastrada no item. Para itens não identificados como SINAPI e sem observação técnica, o memorial exibe alerta documental neutro para complementação antes do protocolo oficial, quando aplicável.',
        styles['TextoPPP']
    ))
    if etapas:
        for idx, (etapa, rows) in enumerate(etapas.items(), start=1):
            story.append(Paragraph(f'<b>{sec_etapas}.{idx} {_html(etapa)}</b>', styles['TextoPPP']))
            story.append(Spacer(1, 3))
            for r in rows:
                _add_item_memorial(story, r, styles)
    else:
        story.append(Paragraph('Nenhuma etapa lançada no orçamento.', styles['TextoPPP']))

    _add_secao(story, sec, 'Critérios de medição e pagamento', styles); sec += 1
    story.append(Paragraph(
        'A medição deverá observar as quantidades executadas, percentuais aprovados no cronograma físico-financeiro, documentos de aceite, registros de campo, relatórios fotográficos, ordens de serviço, notas fiscais, atestos e demais evidências técnicas anexadas ao processo. A autorização de pagamento deve observar o contrato, o aditivo, a ordem de serviço ou instrumento equivalente e os procedimentos internos de fiscalização e aceite.',
        styles['TextoPPP']
    ))

    if tem_opex:
        _add_secao(story, sec, 'Tratamento de OPEX', styles); sec += 1
        story.append(Paragraph(
            'Este orçamento possui previsão de OPEX demonstrada separadamente. O OPEX deve ser tratado como custo recorrente de operação, manutenção, suporte técnico, conservação, inspeção ou serviço continuado, conforme cadastro do item ou da rubrica mensal. Salvo regra contratual específica, o OPEX não compõe o cronograma físico-financeiro do CAPEX de implantação.',
            styles['TextoPPP']
        ))

    _add_secao(story, sec, 'Documentos complementares e anexos', styles); sec += 1
    if anexos:
        data_anexos = [['Tipo', 'Arquivo / referência', 'Observação']]
        for row in anexos:
            nome = _limpar_nome_texto(texto(row,'nome_arquivo') or Path(str(texto(row,'caminho') or '')).name or 'Arquivo não informado')
            tipo = _tipo_anexo_padrao(texto(row,'tipo'), nome)
            obs = _limpar_nome_texto(texto(row,'observacao') or 'Sem observação cadastrada')
            data_anexos.append([tipo, nome, obs])
        story.append(_table_memorial(data_anexos, [115, 205, 170], styles, font_size=6.7))
        story.append(Spacer(1, 6))
    else:
        story.append(Paragraph('Não foram identificados anexos vinculados ao orçamento até a data de geração deste memorial.', styles['TextoPPP']))

    obs_gerais = _limpar_nome_texto(texto(orc,'observacoes') or texto(orc,'texto_proposta') or '').strip()
    if obs_gerais:
        _add_observacoes_gerais_estruturadas(story, sec, obs_gerais, styles); sec += 1

    doc.build(story); return arquivo

def _normalizar_etapa_nome(nome):
    return str(nome or '').strip().lower()


def _mapa_valores_etapas(itens):
    """Calcula o valor distribuível por etapa para o cronograma.

    Regra PPP: o cronograma físico-financeiro deve distribuir o valor da etapa
    pelo percentual de cada parcela. Quando houver itens CAPEX, prioriza CAPEX;
    se não houver marcação CAPEX, usa todos os itens lançados.
    """
    por_id_capex = defaultdict(float); por_nome_capex = defaultdict(float)
    por_id_todos = defaultdict(float); por_nome_todos = defaultdict(float)
    for item in itens:
        valor = total_item(item)
        etapa_id = texto(item, 'etapa_id')
        etapa_nome = texto(item, 'etapa') or texto(item, 'etapa_nome') or 'Etapa não informada'
        nome_key = _normalizar_etapa_nome(etapa_nome)
        if etapa_id:
            por_id_todos[str(etapa_id)] += valor
        por_nome_todos[nome_key] += valor
        if str(texto(item, 'capex_opex')).upper() == 'CAPEX':
            if etapa_id:
                por_id_capex[str(etapa_id)] += valor
            por_nome_capex[nome_key] += valor
    # Usa CAPEX quando houver CAPEX cadastrado; caso contrário usa todos os itens.
    if sum(por_nome_capex.values()) > 0:
        return por_id_capex, por_nome_capex
    return por_id_todos, por_nome_todos


def _valor_etapa_cronograma(info, etapa_id, etapa_nome, total_por_id, total_por_nome):
    """Retorna o valor da etapa para distribuir em P1..P12.

    Preferência: total calculado pelos itens da etapa. Fallback: soma de valores
    já gravados na tabela de cronograma, quando existir.
    """
    if etapa_id and str(etapa_id) in total_por_id and total_por_id[str(etapa_id)] > 0:
        return total_por_id[str(etapa_id)]
    nome_key = _normalizar_etapa_nome(etapa_nome)
    if nome_key in total_por_nome and total_por_nome[nome_key] > 0:
        return total_por_nome[nome_key]
    return sum(info.get('valores_informados', {}).values())


def gerar_pdf_cronograma(orcamento_id):
    orc = buscar_orcamento(orcamento_id)
    cron = buscar_cronograma(orcamento_id)
    itens = buscar_itens(orcamento_id)

    if not orc:
        raise Exception('Orçamento não localizado.')

    pasta = pasta_orcamento(texto(orc, 'codigo')) / '01_Cliente_Final'
    pasta.mkdir(parents=True, exist_ok=True)

    arquivo = pasta / '04_Cronograma_Fisico_Financeiro_PPP.pdf'
    styles = _styles()

    doc = SimpleDocTemplate(
        str(arquivo),
        pagesize=landscape(A3),
        rightMargin=18,
        leftMargin=18,
        topMargin=16,
        bottomMargin=16
    )

    story = []
    _cabecalho_documento(story, orc, 'CRONOGRAMA FÍSICO-FINANCEIRO - PPP', styles)

    info = Table([
        [
            Paragraph(f"<b>Código:</b> {texto(orc,'codigo')}", styles['TextoPequeno']),
            Paragraph(f"<b>Cliente:</b> {texto(orc,'cliente_nome')}", styles['TextoPequeno']),
            Paragraph(f"<b>Município/UF:</b> {texto(orc,'cliente_municipio') or texto(orc,'municipio')}/{texto(orc,'cliente_uf') or texto(orc,'empresa_uf')}", styles['TextoPequeno']),
            Paragraph(f"<b>Segmento:</b> {texto(orc,'segmento_nome')}", styles['TextoPequeno']),
        ],
        [
            Paragraph(f"<b>Contrato:</b> {texto(orc,'contrato_numero')} - {texto(orc,'contrato_nome')}", styles['TextoPequeno']),
            Paragraph(f"<b>Status:</b> {texto(orc,'status')}", styles['TextoPequeno']),
            Paragraph(f"<b>Validade:</b> {texto(orc,'validade')}", styles['TextoPequeno']),
            Paragraph(f"<b>Emissão:</b> {datetime.now().strftime('%d/%m/%Y')}", styles['TextoPequeno']),
        ]
    ], colWidths=[250, 260, 230, 220])

    info.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.25, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))

    story += [info, Spacer(1, 6)]

    objeto = texto(orc, 'objeto') or texto(orc, 'contrato_objeto') or 'Objeto conforme cadastro do orçamento.'
    story.append(Paragraph(f"<b>Objeto:</b> {objeto}", styles['TextoPequeno']))
    story.append(Paragraph(
        "Cronograma físico-financeiro por parcelas efetivamente preenchidas. "
        "A coluna (%) indica a participação de cada etapa no valor global do CAPEX. "
        "As parcelas apresentam valor previsto (R$) e percentual correspondente, com acumulado financeiro ao final.",
        styles['TextoPequeno']
    ))
    story.append(Spacer(1, 6))

    total_por_id, total_por_nome = _mapa_valores_etapas(itens)
    matriz = OrderedDict()

    for row in cron:
        parcela = int(numero(texto(row, 'parcela')) or 0)
        pct = numero(texto(row, 'percentual'))
        val = numero(texto(row, 'valor'))

        if parcela < 1 or parcela > 12 or (pct <= 0 and val <= 0):
            continue

        etapa_id = texto(row, 'etapa_id')
        etapa = (
            texto(row, 'etapa_nome')
            or texto(row, 'etapa')
            or texto(row, 'descricao')
            or (f"Etapa {etapa_id}" if etapa_id else 'Cronograma')
        )

        chave = (str(etapa_id or ''), str(etapa or 'Cronograma'))

        if chave not in matriz:
            matriz[chave] = {
                'etapa_id': etapa_id,
                'nome': etapa,
                'pct': defaultdict(float),
                'valores_informados': defaultdict(float),
            }

        matriz[chave]['pct'][parcela] += pct
        matriz[chave]['valores_informados'][parcela] += val

    parcelas = sorted({
        p
        for info_m in matriz.values()
        for p in set(list(info_m['pct'].keys()) + list(info_m['valores_informados'].keys()))
        if info_m['pct'].get(p, 0) > 0 or info_m['valores_informados'].get(p, 0) > 0
    })

    if not parcelas:
        parcelas = [1]

    linhas = []
    total_geral = 0.0

    for idx, (_, info_m) in enumerate(matriz.items(), start=1):
        etapa = info_m['nome'] or 'Cronograma'
        valor_etapa = _valor_etapa_cronograma(
            info_m,
            info_m.get('etapa_id'),
            etapa,
            total_por_id,
            total_por_nome
        )

        if valor_etapa <= 0:
            valor_etapa = sum(info_m['valores_informados'].values())

        total_geral += valor_etapa
        linhas.append((idx, info_m, etapa, valor_etapa))

    if total_geral <= 0:
        total_geral = sum(total_por_nome.values()) or 0.0

    header1 = ['Item', 'Descrição dos Serviços', 'Valor Total R$', '(%)']
    for p in parcelas:
        header1 += [f'Parcela {p}', '']
    header1 += ['Acum. %']

    header2 = ['', '', '', '']
    for _ in parcelas:
        header2 += ['R$', '%']
    header2 += ['']

    data = [header1, header2]
    total_parcela = defaultdict(float)

    for idx, info_m, etapa, valor_etapa in linhas:
        participacao = (valor_etapa / total_geral * 100) if total_geral else 0

        linha = [
            str(idx),
            Paragraph(str(etapa), styles['TextoTabela']),
            moeda(valor_etapa),
            f"{participacao:.2f}%".replace('.', ',')
        ]

        pct_acumulado_etapa = 0.0

        for p in parcelas:
            pct = info_m['pct'].get(p, 0)
            valor = valor_etapa * pct / 100.0

            if valor <= 0:
                valor = info_m['valores_informados'].get(p, 0)
                pct = (valor / valor_etapa * 100) if valor_etapa else 0

            pct_acumulado_etapa += pct
            total_parcela[p] += valor

            linha += [
                moeda(valor) if valor else '-',
                f"{pct:.2f}%".replace('.', ',') if pct else '0,00%'
            ]

        linha += [f"{pct_acumulado_etapa:.2f}%".replace('.', ',')]
        data.append(linha)

    if len(data) == 2:
        data.append(
            ['-', 'Cronograma ainda não preenchido no sistema.', '-', '-']
            + ['-', '-'] * len(parcelas)
            + ['-']
        )

    valor_global = ['TOTAL POR PARCELA', '', moeda(total_geral), '100,00%']
    perc_acumulada = ['PERCENTAGEM ACUMULADA', '', '', '']
    custo_acumulado = ['VALOR ACUMULADO', '', '', '']

    acumulado = 0.0

    for p in parcelas:
        valor = total_parcela.get(p, 0)
        acumulado += valor
        pct = (valor / total_geral * 100) if total_geral else 0
        pct_acum = (acumulado / total_geral * 100) if total_geral else 0

        valor_global += [moeda(valor), f"{pct:.2f}%".replace('.', ',')]
        perc_acumulada += ['', f"{pct_acum:.2f}%".replace('.', ',')]
        custo_acumulado += [moeda(acumulado), '']

    valor_global += ['']
    perc_acumulada += ['']
    custo_acumulado += ['']

    data += [valor_global, perc_acumulada, custo_acumulado]

    available = 1135
    item_w = 26
    desc_w = 170
    valor_w = 66
    pct_w = 40
    acum_w = 48

    remaining = available - item_w - desc_w - valor_w - pct_w - acum_w
    parcela_pair_w = max(72, min(94, remaining / max(1, len(parcelas))))

    col_widths = [item_w, desc_w, valor_w, pct_w]

    for _ in parcelas:
        col_widths += [parcela_pair_w * 0.60, parcela_pair_w * 0.40]

    col_widths += [acum_w]

    table = Table(data, repeatRows=2, colWidths=col_widths)

    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 1), colors.HexColor(AZUL)),
        ('TEXTCOLOR', (0, 0), (-1, 1), colors.white),
        ('FONTNAME', (0, 0), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 6.1),
        ('LEADING', (0, 0), (-1, -1), 6.8),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#9AA7B8')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 2), (-1, -1), 'RIGHT'),
        ('ALIGN', (1, 2), (1, -1), 'LEFT'),
        ('ROWBACKGROUNDS', (0, 2), (-1, -4), [colors.white, colors.HexColor(CINZA)]),
        ('BACKGROUND', (0, -3), (-1, -1), colors.HexColor(AZUL_CLARO)),
        ('FONTNAME', (0, -3), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -3), (3, -3), colors.HexColor('#D9EAD3')),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]

    col = 4
    for _ in parcelas:
        style_cmds.append(('SPAN', (col, 0), (col + 1, 0)))
        col += 2

    for c in range(0, 4):
        style_cmds.append(('SPAN', (c, 0), (c, 1)))

    style_cmds.append(('SPAN', (len(col_widths)-1, 0), (len(col_widths)-1, 1)))

    footer_start = len(data) - 3
    for r in range(footer_start, len(data)):
        style_cmds.append(('SPAN', (0, r), (3, r)))
        style_cmds.append(('ALIGN', (0, r), (3, r), 'LEFT'))
        style_cmds.append(('LEFTPADDING', (0, r), (3, r), 6))

    table.setStyle(TableStyle(style_cmds))
    story.append(table)

    # Assinatura do cronograma físico-financeiro
    story.append(Spacer(1, 30))
    assinaturas_cronograma = Table([
        ['Responsável Técnico', 'Representante da Contratante / Fiscalização'],
        ['____________________________________', '____________________________________'],
        [texto(orc,'empresa_responsavel_tecnico') or 'Responsável Técnico', texto(orc,'cliente_responsavel') or 'Prefeitura / Fiscalização'],
        [texto(orc,'empresa_registro_tecnico') or '', texto(orc,'cliente_cargo') or '']
    ], colWidths=[300,300])
    assinaturas_cronograma.setStyle(TableStyle([
        ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),7.5),
        ('TOPPADDING',(0,0),(-1,-1),5),
        ('BOTTOMPADDING',(0,0),(-1,-1),5),
    ]))
    story.append(assinaturas_cronograma)

    doc.build(story)
    return arquivo

def gerar_pdf_indice_anexos(orcamento_id):
    orc = buscar_orcamento(orcamento_id); anexos = buscar_anexos(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '01_Cliente_Final'; pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / '05_Indice_de_Anexos.pdf'; styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=landscape(A3), rightMargin=18, leftMargin=18, topMargin=24, bottomMargin=24)
    story = []
    _cabecalho_documento(story, orc, 'ÍNDICE DE ANEXOS', styles)
    story.append(Paragraph(
        'Relação dos documentos complementares vinculados ao orçamento e considerados como suporte documental da proposta.',
        styles['TextoPequeno']
    ))
    story.append(Spacer(1, 6))

    data = [['Nº', 'Tipo', 'Arquivo', 'Observação', 'Data']]
    for idx, row in enumerate(anexos, start=1):
        nome = _limpar_nome_texto(texto(row,'nome_arquivo') or Path(str(texto(row,'caminho') or '')).name)
        tipo = _tipo_anexo_padrao(texto(row,'tipo'), nome)
        data.append([
            f'{idx:02d}',
            _paragraph(tipo, styles),
            Paragraph(_quebrar_texto_longo(nome), styles['TextoTabela']),
            _paragraph(_limpar_nome_texto(texto(row,'observacao')), styles),
            _paragraph(texto(row,'data_upload') or texto(row,'data_inclusao') or '', styles)
        ])
    if len(data) == 1:
        data.append(['-', 'Nenhum anexo registrado.', '', '', ''])

    table = Table(data, repeatRows=1, colWidths=[28,125,240,300,65])
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor(AZUL)),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#9AA7B8')),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),6.6),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, colors.HexColor(CINZA)]),
        ('LEFTPADDING',(0,0),(-1,-1),4),
        ('RIGHTPADDING',(0,0),(-1,-1),4),
        ('TOPPADDING',(0,0),(-1,-1),4),
        ('BOTTOMPADDING',(0,0),(-1,-1),4),
    ]))
    story.append(table)
    doc.build(story); return arquivo



# ============================================================
# PACOTE GERENCIAL INTERNO - USO RESTRITO DA CONCESSIONÁRIA
# ============================================================

TEXTO_GERENCIAL_PADRAO = (
    "A presente composição considera como custo base interno o valor unitário lançado sem aplicação de margem. "
    "Sobre esse valor são aplicados os percentuais internos definidos no orçamento, resultando no preço unitário final "
    "e no preço total final. A versão cliente/prefeitura exibe apenas os preços finais, enquanto esta versão gerencial "
    "demonstra a formação interna do preço para fins de análise econômica, controle de margem e exequibilidade."
)


def pasta_gerencial(orc):
    pasta = pasta_orcamento(texto(orc, 'codigo')) / '00_Revisao_Interna' / 'Gerencial'
    pasta.mkdir(parents=True, exist_ok=True)
    return pasta


def _aviso_gerencial(story, styles, largura=460):
    aviso = Table([
        [Paragraph('<b>DOCUMENTO GERENCIAL INTERNO - USO RESTRITO DA CONCESSIONÁRIA</b>', styles['TextoPPP'])],
        [Paragraph('<b>NÃO ENCAMINHAR AO PODER CONCEDENTE</b>', styles['TextoPPP'])],
    ], colWidths=[largura])
    aviso.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#FFF2CC')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#7A4B00')),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#D6A700')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    aviso.hAlign = 'CENTER'
    story += [aviso, Spacer(1, 8)]


def _calc_formacao(qtd, custo_unit, bdi, margem):
    q = numero(qtd)
    custo_u = numero(custo_unit)
    bdi_p = numero(bdi)
    margem_p = numero(margem)

    custo_total = q * custo_u
    valor_bdi = custo_total * (bdi_p / 100.0)
    base_com_bdi = custo_total + valor_bdi
    valor_margem = base_com_bdi * (margem_p / 100.0)
    preco_total = base_com_bdi + valor_margem
    preco_unit_final = preco_total / q if q else 0.0
    diferenca = preco_total - custo_total
    margem_sobre_custo = (diferenca / custo_total * 100.0) if custo_total else 0.0
    margem_sobre_preco = (diferenca / preco_total * 100.0) if preco_total else 0.0

    return {
        'qtd': q,
        'custo_unit': custo_u,
        'custo_total': custo_total,
        'bdi': bdi_p,
        'valor_bdi': valor_bdi,
        'margem': margem_p,
        'valor_margem': valor_margem,
        'preco_unit_final': preco_unit_final,
        'preco_total': preco_total,
        'diferenca': diferenca,
        'margem_sobre_custo': margem_sobre_custo,
        'margem_sobre_preco': margem_sobre_preco,
    }


def _pct(v):
    return f"{numero(v):,.2f}%".replace(',', 'X').replace('.', ',').replace('X', '.')


def _fmt_qtd(v):
    return f"{numero(v):,.4f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def _style_gerencial(font_size=7.0):
    return TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(AZUL_ESCURO)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#9AA7B8')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), font_size),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor(CINZA)]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(AZUL_CLARO)),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ])


def gerar_pdf_proposta_gerencial_interna(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '01_Proposta_Gerencial_Interna.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=34, bottomMargin=34)
    capex, opex_total, total = _resumos(itens, opex)

    cron = buscar_cronograma(orcamento_id)
    parcelas_ativas = sorted({
        int(numero(texto(r, 'parcela')) or 0)
        for r in cron
        if int(numero(texto(r, 'parcela')) or 0) > 0 and (numero(texto(r, 'percentual')) > 0 or numero(texto(r, 'valor')) > 0)
    })
    qtd_parcelas = len(parcelas_ativas)
    valor_parcela = capex / qtd_parcelas if qtd_parcelas else 0

    story = []
    _cabecalho_documento(story, orc, 'PROPOSTA GERENCIAL INTERNA - FORMAÇÃO DE PREÇO', styles)
    _aviso_gerencial(story, styles, largura=460)
    story += [
        Paragraph('1. Finalidade gerencial', styles['SecaoPPP']),
        Paragraph(TEXTO_GERENCIAL_PADRAO, styles['TextoPPP']),
        Paragraph('2. Identificação do orçamento', styles['SecaoPPP']),
        Paragraph(f"<b>Código:</b> {texto(orc,'codigo')}<br/><b>Cliente:</b> {texto(orc,'cliente_nome')}<br/><b>Objeto:</b> {_limpar_nome_texto(texto(orc,'objeto') or texto(orc,'contrato_objeto'))}", styles['TextoPPP']),
        Paragraph('3. Referência financeira do orçamento', styles['SecaoPPP']),
    ]

    resumo = Table([
        ['Componente', 'Critério', 'Valor'],
        ['CAPEX', 'Implantação / investimento, base do cronograma físico-financeiro', moeda(capex)],
        ['Parcelamento CAPEX', f'{qtd_parcelas} parcelas de {moeda(valor_parcela)}' if qtd_parcelas else 'Conforme cronograma', moeda(capex)],
        ['OPEX mensal previsto', 'Recorrente / operacional, fora do parcelamento CAPEX', moeda(opex_total) + '/mês'],
    ], colWidths=[130, 260, 120])
    resumo.setStyle(_table_style(font_size=8))
    story += [
        resumo,
        Spacer(1, 12),
        Paragraph('4. Critério de segregação CAPEX/OPEX', styles['SecaoPPP']),
        Paragraph('O CAPEX compõe o cronograma físico-financeiro e o parcelamento do investimento. O OPEX é tratado como valor mensal recorrente previsto, vinculado ao início efetivo da operação/manutenção e ao respectivo atesto, não compondo o total parcelado do CAPEX.', styles['TextoPPP']),
        Paragraph('5. Regra de confidencialidade', styles['SecaoPPP']),
        Paragraph('As informações de custo base interno, BDI, margem interna, valor da margem e análise de rentabilidade são de uso restrito da concessionária e não devem ser encaminhadas ao poder concedente, salvo requisição formal específica para fins de auditoria, fiscalização ou exequibilidade.', styles['TextoPPP']),
    ]
    doc.build(story); return arquivo


def gerar_pdf_planilha_gerencial_capex(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '02_Planilha_Gerencial_CAPEX.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=landscape(A2), rightMargin=18, leftMargin=18, topMargin=22, bottomMargin=22)
    story = []
    _cabecalho_documento(story, orc, 'PLANILHA GERENCIAL CAPEX - FORMAÇÃO INTERNA DO PREÇO', styles)
    _aviso_gerencial(story, styles, largura=1500)

    story.append(Paragraph(TEXTO_GERENCIAL_PADRAO, styles['TextoPequeno']))
    story.append(Paragraph(
        'A tabela abaixo demonstra a formação do preço por item CAPEX. As observações internas ficam consolidadas abaixo da tabela para evitar compressão e corte das colunas principais.',
        styles['TextoPequeno']
    ))
    story.append(Spacer(1, 8))

    headers = [
        'Etapa', 'Código', 'Descrição', 'Fonte / referência', 'UN', 'Qtd',
        'Custo unit. sem margem', 'Custo total sem margem', 'BDI %', 'Valor BDI',
        'Margem %', 'Valor margem', 'Preço unit. final', 'Preço total final',
        'Margem s/ custo %', 'Margem s/ preço %'
    ]
    data = [headers]
    observacoes = []
    total_custo = total_bdi = total_margem = total_final = 0.0

    for item in [i for i in itens if str(texto(i, 'capex_opex')).upper() == 'CAPEX']:
        c = _calc_formacao(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem'))
        total_custo += c['custo_total']; total_bdi += c['valor_bdi']; total_margem += c['valor_margem']; total_final += c['preco_total']
        cod = _limpar_nome_texto(texto(item,'codigo'))
        obs = _limpar_nome_texto(texto(item,'observacao') or texto(item,'observacoes'))
        if obs:
            observacoes.append([_paragraph(cod or '-', styles), _paragraph(obs, styles)])
        data.append([
            _paragraph(_limpar_nome_texto(texto(item,'etapa')), styles),
            _paragraph(cod, styles),
            _paragraph(_limpar_nome_texto(texto(item,'descricao')), styles),
            _paragraph(_limpar_nome_texto(texto(item,'fonte') or 'Custo base interno'), styles),
            texto(item,'unidade'),
            _fmt_qtd(c['qtd']),
            moeda(c['custo_unit']),
            moeda(c['custo_total']),
            _pct(c['bdi']),
            moeda(c['valor_bdi']),
            _pct(c['margem']),
            moeda(c['valor_margem']),
            moeda4(c['preco_unit_final']),
            moeda(c['preco_total']),
            _pct(c['margem_sobre_custo']),
            _pct(c['margem_sobre_preco']),
        ])

    data.append([
        'TOTAL CAPEX', '', '', '', '', '',
        '', moeda(total_custo), '', moeda(total_bdi), '', moeda(total_margem),
        '', moeda(total_final), '', ''
    ])

    col_widths = [120, 90, 330, 120, 42, 62, 105, 115, 55, 95, 65, 98, 105, 115, 85, 85]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.hAlign = 'CENTER'
    table.setStyle(_style_gerencial(font_size=7.0))
    story.append(table)

    story.append(Spacer(1, 12))
    story.append(Paragraph('Observações internas CAPEX', styles['SecaoPPP']))
    if observacoes:
        obs_table = Table([['Código', 'Observação interna']] + observacoes, repeatRows=1, colWidths=[120, 1360])
        obs_table.hAlign = 'CENTER'
        obs_table.setStyle(_style_gerencial(font_size=7.0))
        story.append(obs_table)
    else:
        story.append(Paragraph('Sem observações internas registradas para os itens CAPEX.', styles['TextoPequeno']))

    doc.build(story); return arquivo


def gerar_pdf_planilha_gerencial_opex(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '03_Planilha_Gerencial_OPEX.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=landscape(A2), rightMargin=18, leftMargin=18, topMargin=22, bottomMargin=22)
    story = []
    _cabecalho_documento(story, orc, 'PLANILHA GERENCIAL OPEX - FORMAÇÃO INTERNA DO PREÇO', styles)
    _aviso_gerencial(story, styles, largura=1320)
    story.append(Paragraph('O OPEX é demonstrado como valor mensal recorrente previsto e não compõe o cronograma físico-financeiro do CAPEX.', styles['TextoPequeno']))
    story.append(Paragraph('A tabela abaixo demonstra os valores principais. As condições e observações ficam abaixo para evitar planilha cortada e texto comprimido.', styles['TextoPequeno']))
    story.append(Spacer(1, 8))

    headers = [
        'Código', 'Categoria', 'Descrição', 'Fonte', 'UN', 'Qtd mensal',
        'Custo unit. sem margem', 'Custo mensal sem margem', 'BDI %', 'Margem %',
        'Valor margem', 'Preço unit. final mensal', 'Total mensal final',
        'Início previsto cobrança', 'Condição de cobrança'
    ]
    data = [headers]
    observacoes = []
    total_custo = total_margem = total_final = 0.0

    opex_itens = [i for i in itens if str(texto(i, 'capex_opex')).upper() != 'CAPEX']
    for item in opex_itens:
        c = _calc_formacao(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem'))
        total_custo += c['custo_total']; total_margem += c['valor_margem']; total_final += c['preco_total']
        cod = _limpar_nome_texto(texto(item,'codigo'))
        obs = _limpar_nome_texto(texto(item,'observacao') or texto(item,'observacoes'))
        if obs:
            observacoes.append([_paragraph(cod or '-', styles), _paragraph(obs, styles)])
        data.append([
            _paragraph(cod, styles),
            _paragraph(_limpar_nome_texto(texto(item,'classificacao') or texto(item,'etapa') or 'OPEX'), styles),
            _paragraph(_limpar_nome_texto(texto(item,'descricao')), styles),
            _paragraph(_limpar_nome_texto(texto(item,'fonte')), styles),
            texto(item,'unidade'),
            _fmt_qtd(c['qtd']),
            moeda(c['custo_unit']),
            moeda(c['custo_total']),
            _pct(c['bdi']),
            _pct(c['margem']),
            moeda(c['valor_margem']),
            moeda4(c['preco_unit_final']),
            moeda(c['preco_total']),
            _paragraph(texto(item,'inicio_cobranca') or '', styles),
            _paragraph('Após início da operação e atesto' if (texto(item,'condicao_cobranca') or '') else '', styles),
        ])

    for row in opex:
        c = _calc_formacao(texto(row,'qtd_mensal'), texto(row,'custo_unit'), texto(row,'bdi'), texto(row,'margem'))
        total_custo += c['custo_total']; total_margem += c['valor_margem']; total_final += c['preco_total']
        cod = _limpar_nome_texto(texto(row,'codigo'))
        obs = _limpar_nome_texto(texto(row,'observacao') or texto(row,'observacoes'))
        if obs:
            observacoes.append([_paragraph(cod or '-', styles), _paragraph(obs, styles)])
        data.append([
            _paragraph(cod, styles),
            _paragraph(_limpar_nome_texto(texto(row,'categoria') or texto(row,'centro_custo') or 'OPEX'), styles),
            _paragraph(_limpar_nome_texto(texto(row,'descricao')), styles),
            _paragraph(_limpar_nome_texto(texto(row,'fonte')), styles),
            texto(row,'unidade'),
            _fmt_qtd(c['qtd']),
            moeda(c['custo_unit']),
            moeda(c['custo_total']),
            _pct(c['bdi']),
            _pct(c['margem']),
            moeda(c['valor_margem']),
            moeda4(c['preco_unit_final']),
            moeda(c['preco_total']),
            _paragraph(texto(row,'inicio_cobranca') or texto(row,'inicio_previsto') or '', styles),
            _paragraph('Após início da operação e atesto', styles),
        ])

    if len(data) == 1:
        data.append(['-', '-', 'Nenhum item OPEX lançado.', '-', '-', '-', '-', '-', '-', '-', '-', '-', '-', '-', '-'])

    data.append(['TOTAL OPEX MENSAL', '', '', '', '', '', '', moeda(total_custo), '', '', moeda(total_margem), '', moeda(total_final), '', ''])

    col_widths = [80, 95, 260, 85, 38, 55, 100, 105, 50, 55, 90, 110, 110, 90, 145]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.hAlign = 'CENTER'
    table.setStyle(_style_gerencial(font_size=7.2))
    story.append(table)

    story.append(Spacer(1, 12))
    story.append(Paragraph('Observações internas OPEX', styles['SecaoPPP']))
    if observacoes:
        obs_table = Table([['Código', 'Observação interna']] + observacoes, repeatRows=1, colWidths=[120, 1360])
        obs_table.hAlign = 'CENTER'
        obs_table.setStyle(_style_gerencial(font_size=7.0))
        story.append(obs_table)
    else:
        story.append(Paragraph('Sem observações internas registradas para os itens OPEX.', styles['TextoPequeno']))

    doc.build(story); return arquivo


def gerar_pdf_memorial_economico_gerencial(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '04_Memorial_Economico_Gerencial.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=34, bottomMargin=34)
    capex, opex_total, total = _resumos(itens, opex)
    story = []
    _cabecalho_documento(story, orc, 'MEMORIAL ECONÔMICO GERENCIAL', styles)
    _aviso_gerencial(story, styles, largura=460)
    story += [
        Paragraph('1. Metodologia de formação do preço', styles['SecaoPPP']),
        Paragraph(TEXTO_GERENCIAL_PADRAO, styles['TextoPPP']),
        Paragraph('2. Fórmula aplicada', styles['SecaoPPP']),
        Paragraph('Preço final = custo base interno × (1 + BDI%) × (1 + margem interna%). O custo base interno corresponde ao campo “Custo Unitário” lançado no sistema, antes da aplicação de margem.', styles['TextoPPP']),
        Paragraph('3. Separação CAPEX/OPEX', styles['SecaoPPP']),
        Paragraph(f'CAPEX final: {moeda(capex)}. OPEX mensal previsto: {moeda(opex_total)}/mês. OPEX não compõe o cronograma físico-financeiro do CAPEX.', styles['TextoPPP']),
        Paragraph('4. Exequibilidade econômica', styles['SecaoPPP']),
        Paragraph('A versão gerencial permite avaliar se a margem aplicada é suficiente para cobrir riscos operacionais, administrativos, tributários, financeiros e obrigações de execução, preservando a rastreabilidade entre o custo base interno e o preço final demonstrado nos documentos externos.', styles['TextoPPP']),
    ]
    doc.build(story); return arquivo


def gerar_pdf_relatorio_formacao_preco_margem(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '05_Relatorio_Formacao_Preco_Margem.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=landscape(A4), rightMargin=18, leftMargin=18, topMargin=24, bottomMargin=24)
    story = []
    _cabecalho_documento(story, orc, 'RELATÓRIO DE FORMAÇÃO DE PREÇO E MARGEM', styles)
    _aviso_gerencial(story, styles, largura=720)

    data = [['Grupo', 'Custo base', 'Valor BDI', 'Valor margem', 'Preço final', 'Diferença base/final', 'Margem s/ custo', 'Margem s/ preço']]
    grupos = OrderedDict()
    for item in itens:
        grupo = 'CAPEX - ' + str(texto(item,'etapa') or 'Etapa não informada') if str(texto(item,'capex_opex')).upper() == 'CAPEX' else 'OPEX - Itens orçamentários'
        grupos.setdefault(grupo, []).append(_calc_formacao(texto(item,'quantidade'), texto(item,'custo_unit'), texto(item,'bdi'), texto(item,'margem')))
    for row in opex:
        grupos.setdefault('OPEX - Recorrente mensal', []).append(_calc_formacao(texto(row,'qtd_mensal'), texto(row,'custo_unit'), texto(row,'bdi'), texto(row,'margem')))

    for grupo, valores in grupos.items():
        custo = sum(v['custo_total'] for v in valores)
        bdi = sum(v['valor_bdi'] for v in valores)
        margem = sum(v['valor_margem'] for v in valores)
        final = sum(v['preco_total'] for v in valores)
        dif = final - custo
        data.append([
            _paragraph(_limpar_nome_texto(grupo), styles),
            moeda(custo),
            moeda(bdi),
            moeda(margem),
            moeda(final),
            moeda(dif),
            _pct((dif / custo * 100) if custo else 0),
            _pct((dif / final * 100) if final else 0),
        ])

    table = Table(data, repeatRows=1, colWidths=[230, 85, 85, 85, 85, 100, 80, 80])
    table.hAlign = 'CENTER'
    table.setStyle(_style_gerencial(font_size=6.8))
    story.append(table)
    doc.build(story); return arquivo


def gerar_pdf_dashboard_gerencial(orcamento_id):
    orc = buscar_orcamento(orcamento_id); itens = buscar_itens(orcamento_id); opex = buscar_opex(orcamento_id)
    if not orc: raise Exception('Orçamento não localizado.')
    pasta = pasta_gerencial(orc)
    arquivo = pasta / '06_Dashboard_Executivo_Gerencial.pdf'
    styles = _styles()
    doc = SimpleDocTemplate(str(arquivo), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=34, bottomMargin=34)
    capex, opex_total, total = _resumos(itens, opex)
    custo_capex = sum(numero(texto(i,'quantidade')) * numero(texto(i,'custo_unit')) for i in itens if str(texto(i,'capex_opex')).upper() == 'CAPEX')
    custo_opex = sum(numero(texto(i,'quantidade')) * numero(texto(i,'custo_unit')) for i in itens if str(texto(i,'capex_opex')).upper() != 'CAPEX') + sum(numero(texto(o,'qtd_mensal')) * numero(texto(o,'custo_unit')) for o in opex)
    margem_capex = capex - custo_capex
    margem_opex = opex_total - custo_opex
    story = []
    _cabecalho_documento(story, orc, 'DASHBOARD EXECUTIVO GERENCIAL', styles)
    _aviso_gerencial(story, styles, largura=460)
    cards = Table([
        ['Indicador', 'Valor'],
        ['CAPEX final', moeda(capex)],
        ['Custo base CAPEX', moeda(custo_capex)],
        ['Diferença CAPEX base/final', moeda(margem_capex)],
        ['OPEX mensal final', moeda(opex_total) + '/mês'],
        ['Custo base OPEX mensal', moeda(custo_opex) + '/mês'],
        ['Diferença OPEX base/final', moeda(margem_opex) + '/mês'],
    ], colWidths=[260, 180])
    cards.setStyle(_table_style(font_size=8))
    story += [cards, Spacer(1, 10), Paragraph('Este painel é destinado à análise interna de margem, exequibilidade e sustentabilidade econômica do orçamento.', styles['TextoPPP'])]
    doc.build(story); return arquivo



def gerar_pacote_revisao(orcamento_id):
    arquivos = [
        gerar_excel_revisao(orcamento_id),
        gerar_pdf_proposta_gerencial_interna(orcamento_id),
        gerar_pdf_planilha_gerencial_capex(orcamento_id),
        gerar_pdf_planilha_gerencial_opex(orcamento_id),
        gerar_pdf_memorial_economico_gerencial(orcamento_id),
        gerar_pdf_relatorio_formacao_preco_margem(orcamento_id),
        gerar_pdf_dashboard_gerencial(orcamento_id),
        gerar_dashboard_html(orcamento_id, interno=True),
    ]
    return arquivos



# EXPORTAÇÃO DOCX / XLSX
def gerar_docx_proposta_final(orcamento_id):
    return gerar_pdf_proposta_final(orcamento_id)

def gerar_docx_planilha_final(orcamento_id):
    return gerar_pdf_planilha_final(orcamento_id)

def gerar_docx_memorial_descritivo(orcamento_id):
    return gerar_pdf_memorial_descritivo(orcamento_id)

def gerar_xlsx_planilha_final(orcamento_id):
    return gerar_excel_revisao(orcamento_id)

def gerar_xlsx_cronograma_final(orcamento_id):
    return gerar_excel_revisao(orcamento_id)

def gerar_pacote_final(orcamento_id, formatos=None):
    formatos = formatos or ["pdf"]

    arquivos = []

    if "pdf" in formatos:
        arquivos.extend([
            gerar_pdf_proposta_final(orcamento_id),
            gerar_pdf_planilha_final(orcamento_id),
            gerar_pdf_memorial_descritivo(orcamento_id),
            gerar_pdf_cronograma(orcamento_id),
            gerar_pdf_indice_anexos(orcamento_id),
        ])

    if "docx" in formatos:
        arquivos.extend([
            gerar_docx_proposta_final(orcamento_id),
            gerar_docx_planilha_final(orcamento_id),
            gerar_docx_memorial_descritivo(orcamento_id),
        ])

    if "xlsx" in formatos:
        arquivos.extend([
            gerar_xlsx_planilha_final(orcamento_id),
            gerar_xlsx_cronograma_final(orcamento_id),
        ])

    if "html" in formatos:
        arquivos.append(gerar_dashboard_html(orcamento_id, interno=False))

    return arquivos


