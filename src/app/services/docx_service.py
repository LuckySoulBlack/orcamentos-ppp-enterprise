from pathlib import Path
from docx import Document
from docx.shared import Cm

def _salvar_docx(titulo, destino, linhas):
    doc = Document()
    doc.add_heading(titulo, level=1)
    for linha in linhas:
        doc.add_paragraph(str(linha))
    doc.save(destino)
    return str(destino)

def gerar_docx_proposta_final(orcamento_id, pasta_saida):
    arq = Path(pasta_saida) / "01_Proposta_Tecnica_Comercial_PPP.docx"
    return _salvar_docx("Proposta Técnica Comercial PPP", arq, [f"Orçamento: {orcamento_id}"])

def gerar_docx_planilha_final(orcamento_id, pasta_saida):
    arq = Path(pasta_saida) / "02_Planilha_Orcamentaria_PPP.docx"
    return _salvar_docx("Planilha Orçamentária PPP", arq, [f"Orçamento: {orcamento_id}"])

def gerar_docx_memorial_descritivo(orcamento_id, pasta_saida):
    arq = Path(pasta_saida) / "03_Memorial_Descritivo_PPP.docx"
    return _salvar_docx("Memorial Descritivo PPP", arq, [f"Orçamento: {orcamento_id}"])

def gerar_docx_cronograma(orcamento_id, pasta_saida):
    arq = Path(pasta_saida) / "04_Cronograma_Fisico_Financeiro_PPP.docx"
    return _salvar_docx("Cronograma Físico Financeiro PPP", arq, [f"Orçamento: {orcamento_id}"])

def gerar_docx_indice_anexos(orcamento_id, pasta_saida):
    arq = Path(pasta_saida) / "05_Indice_de_Anexos.docx"
    return _salvar_docx("Índice de Anexos", arq, [f"Orçamento: {orcamento_id}"])
