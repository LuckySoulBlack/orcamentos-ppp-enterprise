from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

from core.database import execute, conn


UFS = [
    "AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO",
    "MA", "MG", "MS", "MT", "PA", "PB", "PE", "PI", "PR",
    "RJ", "RN", "RO", "RR", "RS", "SC", "SE", "SP", "TO"
]


ABAS_SINAPI_PADRAO = [
    "ISE", "ICD", "ISD",
    "CSE", "CSD", "CCD",
    "Analítico", "Analitico",
    "Analítico com Custo", "Analitico com Custo",
]


def normalizar(valor):
    return (
        str(valor or "")
        .strip()
        .upper()
        .replace("\n", " ")
        .replace("\r", " ")
        .replace("  ", " ")
    )


def parse_num(valor):
    if valor is None or valor == "":
        return None

    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    if texto in ["-", "–", "—"]:
        return None

    texto = texto.replace("R$", "").replace(" ", "")

    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    try:
        return float(texto)
    except Exception:
        return None


def find_col(headers, termos):
    for i, h in enumerate(headers):
        h_norm = normalizar(h)
        for termo in termos:
            if termo.upper() in h_norm:
                return i
    return None


def inferir_competencia(rows):
    for row in rows[:25]:
        valores = [str(v or "").strip() for v in row]
        joined = " ".join(valores).lower()

        if "mês de referência" in joined or "mes de referencia" in joined:
            for v in valores:
                if "/" in v and any(ch.isdigit() for ch in v):
                    return v

        for v in valores:
            if "/" in v and any(ch.isdigit() for ch in v) and len(v) <= 10:
                if "202" in v or "20" in v:
                    return v

    return ""


def detectar_tipo_aba(nome_aba, rows):
    nome = normalizar(nome_aba)
    texto_topo = " ".join(
        " ".join(normalizar(x) for x in row)
        for row in rows[:15]
    )

    if nome in ("CSE", "CSD", "CCD") or "COMPOSI" in texto_topo or "CUSTOS DE COMPOSI" in texto_topo:
        return "COMPOSICAO"

    if nome in ("ISE", "ICD", "ISD") or "INSUMO" in texto_topo:
        return "INSUMO"

    if "ANALIT" in nome:
        return "COMPOSICAO"

    return "INSUMO"


def localizar_linha_principal(rows):
    """
    Localiza a linha onde ficam código, descrição e unidade.
    Funciona para:
    - ISE/ICD/ISD: insumos;
    - CSE/CSD/CCD: composições;
    - abas Analítico / Analítico com Custo quando a estrutura tiver cabeçalho semelhante.
    """
    for idx, row in enumerate(rows):
        joined = " ".join(normalizar(x) for x in row)

        tem_codigo = "CÓDIGO" in joined or "CODIGO" in joined
        tem_descricao = "DESCRI" in joined
        tem_unidade = "UNIDADE" in joined or " UN " in f" {joined} "

        if tem_codigo and tem_descricao and tem_unidade:
            return idx

    return None


def detectar_uf_cols_insumos(headers):
    uf_cols = []

    for idx, h in enumerate(headers):
        h = normalizar(h)
        if h in UFS:
            uf_cols.append((h, idx))

    return uf_cols


def detectar_uf_cols_composicoes(rows, header_idx):
    """
    Abas CSE/CSD/CCD:
    normalmente a linha acima do cabeçalho possui as UFs
    e a linha do cabeçalho possui Custo (R$) / %AS.
    A coluna da UF é a coluna "Custo (R$)" associada à UF.
    """
    uf_cols = []

    for r_idx in range(max(0, header_idx - 4), min(len(rows), header_idx + 2)):
        row = [normalizar(x) for x in rows[r_idx]]
        for idx, h in enumerate(row):
            if h in UFS:
                uf_cols.append((h, idx))

    # fallback: se as UFs estiverem na própria linha de cabeçalho
    if not uf_cols and 0 <= header_idx < len(rows):
        row = [normalizar(x) for x in rows[header_idx]]
        for idx, h in enumerate(row):
            if h in UFS:
                uf_cols.append((h, idx))

    seen = set()
    final = []
    for uf, col in uf_cols:
        key = (uf, col)
        if key not in seen:
            seen.add(key)
            final.append((uf, col))

    return final


def filtrar_ufs(uf_cols, uf_alvo):
    if not uf_alvo or str(uf_alvo).upper() == "TODAS":
        return uf_cols

    uf_alvo = str(uf_alvo).upper().strip()
    return [(uf, col) for uf, col in uf_cols if uf == uf_alvo]


def limpar_base_sinapi():
    with conn() as c:
        c.execute("DELETE FROM sinapi")
        c.commit()


def preparar_indices_sinapi():
    with conn() as c:
        c.execute("CREATE INDEX IF NOT EXISTS idx_sinapi_codigo ON sinapi(codigo)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_sinapi_descricao ON sinapi(descricao)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_sinapi_uf ON sinapi(uf)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_sinapi_tipo ON sinapi(tipo)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_sinapi_fonte ON sinapi(fonte)")
        c.commit()


def ordenar_abas_sinapi(sheetnames):
    abas = []

    for pref in ABAS_SINAPI_PADRAO:
        for nome in sheetnames:
            if nome.strip().upper() == pref.strip().upper() and nome not in abas:
                abas.append(nome)

    for nome in sheetnames:
        if nome not in abas:
            abas.append(nome)

    return abas


def importar_sinapi_excel(
    caminho,
    uf_alvo="SP",
    limpar_base=False,
    incluir_insumos=True,
    incluir_composicoes=True,
    importar_todas_abas=True,
):
    """
    Importador SINAPI Smart PPP.

    Regra atualizada:
    - importa insumos e composições;
    - lê ISE, ISD, ICD, CSE, CSD, CCD e abas analíticas quando existirem;
    - mantém a fonte com o nome da aba para rastreabilidade;
    - permite pesquisar depois por SINAPI/insumo/composição na tela de orçamento.

    Parâmetros:
    - uf_alvo='SP': importa somente SP.
    - uf_alvo='TODAS': importa todas as UFs.
    - limpar_base=True: apaga base SINAPI antes.
    - incluir_insumos=True: importa ISE/ISD/ICD.
    - incluir_composicoes=True: importa CSE/CSD/CCD/Analítico.
    - importar_todas_abas=True: percorre todas as abas da planilha.
    """
    caminho = Path(caminho)

    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    if limpar_base:
        limpar_base_sinapi()

    preparar_indices_sinapi()

    wb = load_workbook(caminho, data_only=True, read_only=True)
    abas = ordenar_abas_sinapi(wb.sheetnames) if importar_todas_abas else ordenar_abas_sinapi(wb.sheetnames)[:6]

    total = 0
    diagnostico = []
    data_atualizacao = datetime.now().strftime("%d/%m/%Y")

    with conn() as c:
        for nome_aba in abas:
            ws = wb[nome_aba]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                continue

            tipo = detectar_tipo_aba(nome_aba, rows)

            if tipo == "INSUMO" and not incluir_insumos:
                diagnostico.append(f"{nome_aba}: ignorada por configuração (insumos desativados)")
                continue

            if tipo == "COMPOSICAO" and not incluir_composicoes:
                diagnostico.append(f"{nome_aba}: ignorada por configuração (composições desativadas)")
                continue

            header_idx = localizar_linha_principal(rows)

            if header_idx is None:
                diagnostico.append(f"{nome_aba}: cabeçalho não localizado")
                continue

            headers = [normalizar(h) for h in rows[header_idx]]

            col_classificacao = find_col(headers, ["CLASSIFICAÇÃO", "CLASSIFICACAO", "GRUPO"])
            col_codigo = find_col(headers, [
                "CÓDIGO DO INSUMO",
                "CODIGO DO INSUMO",
                "CÓDIGO DA COMPOSIÇÃO",
                "CODIGO DA COMPOSICAO",
                "CÓDIGO",
                "CODIGO"
            ])
            col_descricao = find_col(headers, [
                "DESCRIÇÃO DO INSUMO",
                "DESCRICAO DO INSUMO",
                "DESCRIÇÃO DA COMPOSIÇÃO",
                "DESCRICAO DA COMPOSICAO",
                "DESCRIÇÃO",
                "DESCRICAO"
            ])
            col_unidade = find_col(headers, ["UNIDADE", "UN"])

            if tipo == "INSUMO":
                uf_cols = detectar_uf_cols_insumos(headers)
            else:
                uf_cols = detectar_uf_cols_composicoes(rows, header_idx)

            uf_cols = filtrar_ufs(uf_cols, uf_alvo)

            if col_codigo is None or col_descricao is None or col_unidade is None or not uf_cols:
                diagnostico.append(
                    f"{nome_aba}: colunas incompletas "
                    f"(codigo={col_codigo}, descricao={col_descricao}, unidade={col_unidade}, ufs={len(uf_cols)}, tipo={tipo})"
                )
                continue

            competencia = inferir_competencia(rows)
            total_aba = 0

            for row in rows[header_idx + 1:]:
                if col_codigo >= len(row) or col_descricao >= len(row):
                    continue

                codigo = str(row[col_codigo] or "").strip()
                descricao = str(row[col_descricao] or "").strip()

                if not codigo or not descricao:
                    continue

                if codigo.upper() in ["CÓDIGO", "CODIGO", "NAN", "NONE"]:
                    continue

                if "ACESSAR FICHAS" in descricao.upper():
                    continue

                unidade = str(row[col_unidade] or "").strip() if col_unidade < len(row) else ""

                classificacao = ""
                if col_classificacao is not None and col_classificacao < len(row):
                    classificacao = str(row[col_classificacao] or "").strip()

                for uf, col_uf in uf_cols:
                    if col_uf >= len(row):
                        continue

                    custo = parse_num(row[col_uf])

                    if custo is None:
                        continue

                    c.execute("""
                        INSERT INTO sinapi(
                            codigo,
                            descricao,
                            unidade,
                            custo,
                            uf,
                            competencia,
                            tipo,
                            fonte,
                            data_atualizacao,
                            classificacao
                        )
                        VALUES(?,?,?,?,?,?,?,?,?,?)
                    """, (
                        codigo,
                        descricao,
                        unidade,
                        custo,
                        uf,
                        competencia,
                        tipo,
                        f"SINAPI Excel - {nome_aba}",
                        data_atualizacao,
                        classificacao
                    ))

                    total += 1
                    total_aba += 1

            diagnostico.append(f"{nome_aba}: {total_aba} registros importados ({tipo})")

        c.commit()

    if total == 0:
        raise Exception(
            "Nenhum item SINAPI foi importado.\n\n"
            "Diagnóstico:\n" + "\n".join(diagnostico[:50]) +
            "\n\nDica: selecione a UF SP ou TODAS e confirme se a planilha possui abas ISE/ICD/CSE/CSD."
        )

    preparar_indices_sinapi()
    return total, diagnostico
