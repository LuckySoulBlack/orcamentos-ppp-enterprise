from pathlib import Path
from datetime import datetime, timedelta

from PySide6.QtCore import Qt, QDate
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QLineEdit,
    QPushButton, QComboBox, QDateEdit, QGroupBox, QScrollArea, QTableWidget,
    QTableWidgetItem, QHeaderView, QMessageBox, QFrame, QFileDialog,
    QCheckBox, QAbstractItemView
)

from core.database import conn, STATUS_ORCAMENTO, get_pasta_local_relatorios
from widgets.table_utils import configurar_tabela_responsiva, ajustar_tabela_apos_carregar

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except Exception:
    Workbook = None


class RelatoriosView(QWidget):
    """Central de relatórios gerenciais PPP.

    Objetivo da tela:
    - consultar orçamentos por status, contrato, aditivo, segmento, cliente, município e período;
    - consolidar CAPEX/OPEX, itens, anexos, fontes e parcelas ativas;
    - exportar Excel manipulável com abas completas para gestão, diretoria e auditoria interna.
    """

    def __init__(self):
        super().__init__()
        self._build()
        self.carregar_filtros()
        self.carregar_relatorio()

    # ------------------------- UI -------------------------
    def botao(self, texto, cor="#0B3B75"):
        btn = QPushButton(texto)
        btn.setMinimumHeight(40)
        btn.setCursor(Qt.PointingHandCursor)
        btn.setStyleSheet(f"""
            QPushButton {{ background:{cor}; color:white; border:none; border-radius:8px; font-weight:850; padding:8px 14px; }}
            QPushButton:hover {{ background:{cor}; opacity:.92; }}
        """)
        return btn

    def moeda(self, valor):
        try:
            valor = float(valor or 0)
        except Exception:
            valor = 0.0
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def _scroll_page(self):
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        content = QWidget()
        layout = QVBoxLayout(content)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(14)
        scroll.setWidget(content)
        return scroll, layout

    def _build(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(12, 10, 12, 10)
        root.setSpacing(10)

        title_row = QHBoxLayout()
        title = QLabel("Relatórios Gerenciais PPP")
        title.setStyleSheet("font-size:26px;font-weight:950;color:#071A33;")
        subtitle = QLabel("Consolidação executiva de orçamentos, contratos, aditivos, CAPEX/OPEX, status e pendências.")
        subtitle.setStyleSheet("color:#64748B;font-weight:700;")
        title_box = QVBoxLayout()
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        title_row.addLayout(title_box, 1)
        self.btn_atualizar = self.botao("Atualizar relatório", "#2563EB")
        self.btn_exportar = self.botao("Exportar Excel completo", "#16A34A")
        self.btn_limpar = self.botao("Limpar filtros", "#64748B")
        self.btn_atualizar.clicked.connect(self.carregar_relatorio)
        self.btn_exportar.clicked.connect(self.exportar_excel_completo)
        self.btn_limpar.clicked.connect(self.limpar_filtros)
        title_row.addWidget(self.btn_atualizar)
        title_row.addWidget(self.btn_exportar)
        title_row.addWidget(self.btn_limpar)
        root.addLayout(title_row)

        scroll, layout = self._scroll_page()
        root.addWidget(scroll)

        filtros = QGroupBox("Filtros executivos")
        grid = QGridLayout(filtros)
        grid.setHorizontalSpacing(12)
        grid.setVerticalSpacing(8)

        self.f_cliente = QComboBox(); self.f_cliente.addItem("Todos", None)
        self.f_contrato = QComboBox(); self.f_contrato.addItem("Todos", None)
        self.f_aditivo = QComboBox(); self.f_aditivo.addItem("Todos", None)
        self.f_segmento = QComboBox(); self.f_segmento.addItem("Todos", None)
        self.f_status = QComboBox(); self.f_status.addItem("Todos", None); self.f_status.addItems(STATUS_ORCAMENTO)
        self.f_municipio = QLineEdit(); self.f_municipio.setPlaceholderText("Município ou UF")
        self.f_fonte = QComboBox(); self.f_fonte.addItem("Todas", None)
        self.f_periodo_de = QDateEdit(); self.f_periodo_de.setCalendarPopup(True); self.f_periodo_de.setDate(QDate.currentDate().addYears(-1))
        self.f_periodo_ate = QDateEdit(); self.f_periodo_ate.setCalendarPopup(True); self.f_periodo_ate.setDate(QDate.currentDate())
        self.chk_periodo = QCheckBox("Filtrar por data de criação/atualização")
        self.chk_periodo.setChecked(False)

        fields = [
            ("Cliente", self.f_cliente), ("Contrato", self.f_contrato), ("Aditivo", self.f_aditivo),
            ("Segmento", self.f_segmento), ("Status", self.f_status), ("Município/UF", self.f_municipio),
            ("Fonte", self.f_fonte), ("De", self.f_periodo_de), ("Até", self.f_periodo_ate),
        ]
        for i, (label, widget) in enumerate(fields):
            r = (i // 3) * 2
            c = i % 3
            grid.addWidget(QLabel(label), r, c)
            widget.setMinimumHeight(36)
            grid.addWidget(widget, r + 1, c)
        grid.addWidget(self.chk_periodo, 6, 0, 1, 3)
        layout.addWidget(filtros)

        cards = QGroupBox("Indicadores consolidados")
        cards_lay = QHBoxLayout(cards)
        self.card_qtd = self._card("Orçamentos", "0")
        self.card_capex = self._card("CAPEX", "R$ 0,00")
        self.card_opex = self._card("OPEX mensal", "R$ 0,00")
        self.card_enviados = self._card("Enviados", "0")
        self.card_alertas = self._card("Pendências", "0")
        for card in [self.card_qtd, self.card_capex, self.card_opex, self.card_enviados, self.card_alertas]:
            cards_lay.addWidget(card)
        layout.addWidget(cards)

        self.tabela = QTableWidget(0, 16)
        self.tabela.setHorizontalHeaderLabels([
            "ID", "Código", "Status", "Cliente", "Município", "Contrato", "Aditivo", "Segmento",
            "Objeto", "CAPEX", "OPEX mensal", "Itens CAPEX", "Itens OPEX", "Anexos", "Fontes", "Parcelas ativas"
        ])
        configurar_tabela_responsiva(self.tabela, stretch_columns=[8], hidden_columns=[0])
        self.tabela.setMinimumHeight(410)
        layout.addWidget(self.tabela)

        pend = QGroupBox("Pendências e alertas internos")
        pend_lay = QVBoxLayout(pend)
        self.tabela_pendencias = QTableWidget(0, 5)
        self.tabela_pendencias.setHorizontalHeaderLabels(["Código", "Status", "Tipo", "Descrição", "Prioridade"])
        configurar_tabela_responsiva(self.tabela_pendencias, stretch_columns=[3])
        self.tabela_pendencias.setMinimumHeight(220)
        pend_lay.addWidget(self.tabela_pendencias)
        layout.addWidget(pend)

        layout.addStretch()

    def _card(self, titulo, valor):
        lbl = QLabel(f"{titulo}\n{valor}")
        lbl.setAlignment(Qt.AlignCenter)
        lbl.setMinimumHeight(78)
        lbl.setStyleSheet("""
            QLabel { background:white; border:1px solid #DBE4EF; border-radius:12px; padding:10px;
            font-weight:900; color:#0B3B75; }
        """)
        return lbl

    def _set_card(self, card, titulo, valor):
        card.setText(f"{titulo}\n{valor}")

    # ------------------------- DATA -------------------------
    def carregar_filtros(self):
        with conn() as c:
            clientes = c.execute("SELECT id,nome,uf FROM clientes ORDER BY nome").fetchall()
            contratos = c.execute("SELECT id,COALESCE(numero,'') numero,COALESCE(nome,'Contrato') nome FROM contratos ORDER BY id DESC").fetchall()
            aditivos = c.execute("SELECT id,COALESCE(numero_aditivo,'') numero,COALESCE(objeto,'') objeto FROM aditivos ORDER BY id DESC").fetchall()
            segmentos = c.execute("SELECT id,nome FROM segmentos ORDER BY nome").fetchall()
            fontes = c.execute("""
                SELECT DISTINCT fonte FROM (
                    SELECT fonte FROM itens WHERE fonte IS NOT NULL AND TRIM(fonte) <> ''
                    UNION
                    SELECT fonte FROM opex WHERE fonte IS NOT NULL AND TRIM(fonte) <> ''
                ) ORDER BY fonte
            """).fetchall()

        def refill(combo, rows, formatter):
            atual = combo.currentData()
            combo.clear(); combo.addItem("Todos", None)
            for r in rows:
                combo.addItem(formatter(r), r["id"] if "id" in r.keys() else r[0])
            idx = combo.findData(atual)
            if idx >= 0:
                combo.setCurrentIndex(idx)

        refill(self.f_cliente, clientes, lambda r: f"{r['nome']} - {r['uf'] or ''}")
        refill(self.f_contrato, contratos, lambda r: f"{r['numero'] or r['id']} | {r['nome']}")
        refill(self.f_aditivo, aditivos, lambda r: f"{r['numero'] or r['id']} | {r['objeto'][:70]}")
        refill(self.f_segmento, segmentos, lambda r: r["nome"])
        self.f_fonte.clear(); self.f_fonte.addItem("Todas", None)
        for r in fontes:
            self.f_fonte.addItem(r["fonte"], r["fonte"])

    def _where_params(self):
        where = []
        params = []
        if self.f_cliente.currentData():
            where.append("o.cliente_id=?"); params.append(self.f_cliente.currentData())
        if self.f_contrato.currentData():
            where.append("o.contrato_id=?"); params.append(self.f_contrato.currentData())
        if self.f_aditivo.currentData():
            where.append("o.aditivo_id=?"); params.append(self.f_aditivo.currentData())
        if self.f_segmento.currentData():
            where.append("o.segmento_id=?"); params.append(self.f_segmento.currentData())
        if self.f_status.currentText() and self.f_status.currentText() != "Todos":
            where.append("o.status=?"); params.append(self.f_status.currentText())
        municipio = self.f_municipio.text().strip()
        if municipio:
            where.append("UPPER(COALESCE(o.municipio,'')) LIKE UPPER(?)"); params.append(f"%{municipio}%")
        if self.chk_periodo.isChecked():
            de = self.f_periodo_de.date().toString("yyyy-MM-dd")
            ate = self.f_periodo_ate.date().addDays(1).toString("yyyy-MM-dd")
            where.append("date(COALESCE(o.data_atualizacao,o.data_criacao)) >= date(?)"); params.append(de)
            where.append("date(COALESCE(o.data_atualizacao,o.data_criacao)) < date(?)"); params.append(ate)
        fonte = self.f_fonte.currentData()
        if fonte:
            where.append("""
                EXISTS(SELECT 1 FROM itens i WHERE i.orcamento_id=o.id AND i.fonte=?)
                OR EXISTS(SELECT 1 FROM opex ox WHERE ox.orcamento_id=o.id AND ox.fonte=?)
            """)
            params.extend([fonte, fonte])
        return (" WHERE " + " AND ".join(f"({w})" for w in where)) if where else "", params

    def buscar_orcamentos(self):
        where, params = self._where_params()
        sql = f"""
            SELECT
                o.id, o.codigo, o.status, o.objeto, o.municipio, o.validade, o.data_criacao, o.data_atualizacao,
                cl.nome AS cliente_nome,
                co.numero AS contrato_numero, co.nome AS contrato_nome,
                ad.numero_aditivo AS aditivo_numero, ad.objeto AS aditivo_objeto,
                s.nome AS segmento_nome,
                COALESCE((SELECT SUM(i.quantidade*i.custo_unit*(1+COALESCE(i.bdi,0)/100.0)*(1+COALESCE(i.margem,0)/100.0)) FROM itens i WHERE i.orcamento_id=o.id AND i.capex_opex='CAPEX'),0) AS capex_total,
                COALESCE((SELECT SUM(ox.qtd_mensal*ox.custo_unit*(1+COALESCE(ox.bdi,0)/100.0)*(1+COALESCE(ox.margem,0)/100.0)) FROM opex ox WHERE ox.orcamento_id=o.id),0) AS opex_total,
                COALESCE((SELECT COUNT(*) FROM itens i WHERE i.orcamento_id=o.id AND i.capex_opex='CAPEX'),0) AS itens_capex,
                COALESCE((SELECT COUNT(*) FROM opex ox WHERE ox.orcamento_id=o.id),0) AS itens_opex,
                COALESCE((SELECT COUNT(*) FROM documentos d WHERE d.orcamento_id=o.id),0) AS anexos,
                COALESCE((SELECT COUNT(DISTINCT p.parcela) FROM cronograma_parcelas p WHERE p.orcamento_id=o.id AND (COALESCE(p.percentual,0)>0 OR COALESCE(p.valor,0)>0)),0) AS parcelas_ativas
            FROM orcamentos o
            LEFT JOIN clientes cl ON cl.id=o.cliente_id
            LEFT JOIN contratos co ON co.id=o.contrato_id
            LEFT JOIN aditivos ad ON ad.id=o.aditivo_id
            LEFT JOIN segmentos s ON s.id=o.segmento_id
            {where}
            ORDER BY o.id DESC
        """
        with conn() as c:
            rows = c.execute(sql, params).fetchall()
        return rows

    def fontes_orcamento(self, orcamento_id):
        with conn() as c:
            rows = c.execute("""
                SELECT DISTINCT fonte FROM (
                    SELECT fonte FROM itens WHERE orcamento_id=?AND fonte IS NOT NULL AND TRIM(fonte) <> ''
                    UNION
                    SELECT fonte FROM opex WHERE orcamento_id=?AND fonte IS NOT NULL AND TRIM(fonte) <> ''
                ) ORDER BY fonte
            """, (orcamento_id, orcamento_id)).fetchall()
        return " • ".join([r["fonte"] for r in rows])

    def carregar_relatorio(self):
        rows = self.buscar_orcamentos()
        self.tabela.setRowCount(len(rows))
        total_capex = total_opex = 0.0
        enviados = 0
        pendencias = []
        for r_idx, r in enumerate(rows):
            total_capex += float(r["capex_total"] or 0)
            total_opex += float(r["opex_total"] or 0)
            if (r["status"] or "").upper() == "ENVIADO":
                enviados += 1
            contrato = ""
            if r["contrato_numero"] or r["contrato_nome"]:
                contrato = f"{r['contrato_numero'] or ''} | {r['contrato_nome'] or ''}".strip(" |")
            aditivo = ""
            if r["aditivo_numero"] or r["aditivo_objeto"]:
                aditivo = f"{r['aditivo_numero'] or ''} | {(r['aditivo_objeto'] or '')[:80]}".strip(" |")
            fontes = self.fontes_orcamento(r["id"])
            values = [
                r["id"], r["codigo"], r["status"], r["cliente_nome"], r["municipio"], contrato, aditivo,
                r["segmento_nome"], r["objeto"], self.moeda(r["capex_total"]), self.moeda(r["opex_total"]),
                r["itens_capex"], r["itens_opex"], r["anexos"], fontes, r["parcelas_ativas"]
            ]
            for c_idx, val in enumerate(values):
                item = QTableWidgetItem(str(val or ""))
                if c_idx in (9, 10, 11, 12, 13, 15):
                    item.setTextAlignment(Qt.AlignCenter)
                self.tabela.setItem(r_idx, c_idx, item)
            pendencias.extend(self._pendencias_orcamento(r, fontes))
        ajustar_tabela_apos_carregar(self.tabela)

        self._set_card(self.card_qtd, "Orçamentos", str(len(rows)))
        self._set_card(self.card_capex, "CAPEX", self.moeda(total_capex))
        self._set_card(self.card_opex, "OPEX mensal", self.moeda(total_opex))
        self._set_card(self.card_enviados, "Enviados", str(enviados))
        self._set_card(self.card_alertas, "Pendências", str(len(pendencias)))
        self._carregar_pendencias(pendencias)

    def _pendencias_orcamento(self, r, fontes):
        pend = []
        codigo = r["codigo"] or f"ID {r['id']}"
        status = r["status"] or ""
        if not r["contrato_numero"] and not r["contrato_nome"]:
            pend.append([codigo, status, "Contrato", "Orçamento sem contrato PPP vinculado.", "Alta"])
        if not r["segmento_nome"]:
            pend.append([codigo, status, "Segmento", "Orçamento sem segmento vinculado.", "Média"])
        if int(r["itens_capex"] or 0) == 0 and float(r["opex_total"] or 0) == 0:
            pend.append([codigo, status, "Itens", "Orçamento sem CAPEX e sem OPEX lançado.", "Alta"])
        if not fontes:
            pend.append([codigo, status, "Fonte", "Nenhuma fonte cadastrada nos itens CAPEX/OPEX.", "Alta"])
        if int(r["anexos"] or 0) == 0:
            pend.append([codigo, status, "Anexos", "Nenhum anexo vinculado ao orçamento.", "Média"])
        if float(r["capex_total"] or 0) > 0 and int(r["parcelas_ativas"] or 0) == 0:
            pend.append([codigo, status, "Cronograma", "CAPEX com valor, mas sem parcelas ativas no cronograma.", "Alta"])
        return pend

    def _carregar_pendencias(self, pendencias):
        self.tabela_pendencias.setRowCount(len(pendencias))
        for r, row in enumerate(pendencias):
            for c, val in enumerate(row):
                item = QTableWidgetItem(str(val or ""))
                if c == 4:
                    item.setTextAlignment(Qt.AlignCenter)
                self.tabela_pendencias.setItem(r, c, item)
        ajustar_tabela_apos_carregar(self.tabela_pendencias)

    def limpar_filtros(self):
        for combo in [self.f_cliente, self.f_contrato, self.f_aditivo, self.f_segmento, self.f_status, self.f_fonte]:
            combo.setCurrentIndex(0)
        self.f_municipio.clear()
        self.chk_periodo.setChecked(False)
        self.carregar_relatorio()

    # ------------------------- EXPORTAÇÃO EXCEL -------------------------
    def exportar_excel_completo(self):
        if Workbook is None:
            QMessageBox.critical(self, "Excel", "Biblioteca openpyxl não encontrada. Instale/openpyxl no ambiente do sistema.")
            return

        pasta_padrao = get_pasta_local_relatorios() / "Relatorios_Gerenciais"
        pasta_padrao.mkdir(parents=True, exist_ok=True)
        nome = f"Relatorio_Gerencial_PPP_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        caminho_padrao = pasta_padrao / nome
        path, _ = QFileDialog.getSaveFileName(self, "Salvar Excel gerencial", str(caminho_padrao), "Excel (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        rows = self.buscar_orcamentos()
        wb = Workbook()
        ws_dash = wb.active
        ws_dash.title = "Dashboard"
        self._montar_aba_dashboard(ws_dash, rows)
        self._montar_aba_orcamentos(wb.create_sheet("Orçamentos"), rows)
        self._montar_aba_capex(wb.create_sheet("Itens CAPEX"), rows)
        self._montar_aba_opex(wb.create_sheet("Itens OPEX"), rows)
        self._montar_aba_cronograma(wb.create_sheet("Cronograma"), rows)
        self._montar_aba_contratos(wb.create_sheet("Contratos"))
        self._montar_aba_aditivos(wb.create_sheet("Aditivos"))
        self._montar_aba_anexos(wb.create_sheet("Anexos"), rows)
        self._montar_aba_pendencias(wb.create_sheet("Pendências"), rows)

        for ws in wb.worksheets:
            self._formatar_planilha(ws)
        wb.save(path)
        QMessageBox.information(self, "Excel gerencial", f"Relatório exportado com sucesso:\n{path}")

    def _ids_filtrados(self, rows):
        return [int(r["id"]) for r in rows]

    def _formatar_header(self, ws, row=1):
        fill = PatternFill("solid", fgColor="0B3B75")
        font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="DBE4EF")
        for cell in ws[row]:
            cell.fill = fill
            cell.font = font
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _formatar_planilha(self, ws):
        thin = Side(style="thin", color="DBE4EF")
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in ws.columns:
            max_len = 8
            col_letter = get_column_letter(col[0].column)
            for cell in col[:200]:
                max_len = max(max_len, len(str(cell.value or "")) + 2)
            ws.column_dimensions[col_letter].width = min(max_len, 42)
        ws.freeze_panes = "A2"
        try:
            ws.auto_filter.ref = ws.dimensions
        except Exception:
            pass

    def _montar_aba_dashboard(self, ws, rows):
        total_capex = sum(float(r["capex_total"] or 0) for r in rows)
        total_opex = sum(float(r["opex_total"] or 0) for r in rows)
        enviados = sum(1 for r in rows if (r["status"] or "").upper() == "ENVIADO")
        aprovados = sum(1 for r in rows if (r["status"] or "").upper() == "APROVADO")
        ws.append(["Indicador", "Valor"])
        ws.append(["Orçamentos filtrados", len(rows)])
        ws.append(["CAPEX consolidado", total_capex])
        ws.append(["OPEX mensal consolidado", total_opex])
        ws.append(["Enviados", enviados])
        ws.append(["Aprovados", aprovados])
        ws.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
        self._formatar_header(ws, 1)
        for r in range(3, 5):
            ws[f"B{r}"].number_format = 'R$ #,##0.00'

    def _montar_aba_orcamentos(self, ws, rows):
        ws.append(["ID", "Código", "Status", "Cliente", "Município", "Contrato", "Aditivo", "Segmento", "Objeto", "CAPEX", "OPEX mensal", "Itens CAPEX", "Itens OPEX", "Anexos", "Fontes", "Parcelas ativas", "Validade", "Data criação", "Atualização"])
        for r in rows:
            contrato = f"{r['contrato_numero'] or ''} | {r['contrato_nome'] or ''}".strip(" |")
            aditivo = f"{r['aditivo_numero'] or ''} | {r['aditivo_objeto'] or ''}".strip(" |")
            ws.append([r["id"], r["codigo"], r["status"], r["cliente_nome"], r["municipio"], contrato, aditivo, r["segmento_nome"], r["objeto"], float(r["capex_total"] or 0), float(r["opex_total"] or 0), r["itens_capex"], r["itens_opex"], r["anexos"], self.fontes_orcamento(r["id"]), r["parcelas_ativas"], r["validade"], r["data_criacao"], r["data_atualizacao"]])
        self._formatar_header(ws, 1)
        for col in ["J", "K"]:
            for cell in ws[col][1:]:
                cell.number_format = 'R$ #,##0.00'

    def _montar_aba_capex(self, ws, rows):
        ids = self._ids_filtrados(rows)
        ws.append(["Orçamento", "Status", "Cliente", "Segmento", "Etapa", "Fonte", "Código", "Descrição", "UN", "Qtd", "Custo unit.", "BDI %", "Margem %", "Preço total", "Classificação", "Observação"])
        if not ids:
            self._formatar_header(ws, 1); return
        ph = ",".join("?" for _ in ids)
        with conn() as c:
            dados = c.execute(f"""
                SELECT o.codigo, o.status, cl.nome cliente, s.nome segmento, COALESCE(e.nome,i.etapa) etapa, i.*
                FROM itens i
                JOIN orcamentos o ON o.id=i.orcamento_id
                LEFT JOIN clientes cl ON cl.id=o.cliente_id
                LEFT JOIN segmentos s ON s.id=o.segmento_id
                LEFT JOIN etapas_orcamento e ON e.id=i.etapa_id
                WHERE i.orcamento_id IN ({ph}) AND i.capex_opex='CAPEX'
                ORDER BY o.id DESC, COALESCE(e.ordem,9999), i.id
            """, ids).fetchall()
        for r in dados:
            total = float(r["quantidade"] or 0) * float(r["custo_unit"] or 0) * (1 + float(r["bdi"] or 0)/100) * (1 + float(r["margem"] or 0)/100)
            ws.append([r["codigo"], r["status"], r["cliente"], r["segmento"], r["etapa"], r["fonte"], r["codigo"], r["descricao"], r["unidade"], r["quantidade"], r["custo_unit"], r["bdi"], r["margem"], total, r["classificacao"], r["observacao"]])
        self._formatar_header(ws, 1)

    def _montar_aba_opex(self, ws, rows):
        ids = self._ids_filtrados(rows)
        ws.append(["Orçamento", "Status", "Cliente", "Segmento", "Categoria", "Fonte", "Código", "Descrição", "UN", "Qtd mensal", "Custo unit.", "BDI %", "Margem %", "Total mensal", "Observação"])
        if not ids:
            self._formatar_header(ws, 1); return
        ph = ",".join("?" for _ in ids)
        with conn() as c:
            dados = c.execute(f"""
                SELECT o.codigo, o.status, cl.nome cliente, s.nome segmento, ox.*
                FROM opex ox
                JOIN orcamentos o ON o.id=ox.orcamento_id
                LEFT JOIN clientes cl ON cl.id=o.cliente_id
                LEFT JOIN segmentos s ON s.id=o.segmento_id
                WHERE ox.orcamento_id IN ({ph})
                ORDER BY o.id DESC, ox.id
            """, ids).fetchall()
        for r in dados:
            total = float(r["qtd_mensal"] or 0) * float(r["custo_unit"] or 0) * (1 + float(r["bdi"] or 0)/100) * (1 + float(r["margem"] or 0)/100)
            ws.append([r["codigo"], r["status"], r["cliente"], r["segmento"], r["categoria"], r["fonte"], r["codigo"], r["descricao"], r["unidade"], r["qtd_mensal"], r["custo_unit"], r["bdi"], r["margem"], total, r["observacao"]])
        self._formatar_header(ws, 1)

    def _montar_aba_cronograma(self, ws, rows):
        ids = self._ids_filtrados(rows)
        ws.append(["Orçamento", "Status", "Etapa", "Parcela", "Percentual", "Valor"])
        if not ids:
            self._formatar_header(ws, 1); return
        ph = ",".join("?" for _ in ids)
        with conn() as c:
            dados = c.execute(f"""
                SELECT o.codigo, o.status, e.nome etapa, p.parcela, p.percentual, p.valor
                FROM cronograma_parcelas p
                JOIN orcamentos o ON o.id=p.orcamento_id
                LEFT JOIN etapas_orcamento e ON e.id=p.etapa_id
                WHERE p.orcamento_id IN ({ph}) AND (COALESCE(p.percentual,0)>0 OR COALESCE(p.valor,0)>0)
                ORDER BY o.id DESC, COALESCE(e.ordem,9999), p.parcela
            """, ids).fetchall()
        for r in dados:
            ws.append([r["codigo"], r["status"], r["etapa"], r["parcela"], r["percentual"], r["valor"]])
        self._formatar_header(ws, 1)

    def _montar_aba_contratos(self, ws):
        ws.append(["ID", "Cliente", "Nome", "Número", "Objeto", "Município", "Valor global", "Início", "Fim", "Status", "Índice reajuste", "Observação"])
        with conn() as c:
            dados = c.execute("""
                SELECT co.*, cl.nome cliente
                FROM contratos co LEFT JOIN clientes cl ON cl.id=co.cliente_id
                ORDER BY co.id DESC
            """).fetchall()
        for r in dados:
            ws.append([r["id"], r["cliente"], r["nome"], r["numero"], r["objeto"], r["municipio"], r["valor_global"], r["inicio"], r["fim"], r["status"], r["indice_reajuste"], r["observacao"]])
        self._formatar_header(ws, 1)

    def _montar_aba_aditivos(self, ws):
        ws.append(["ID", "Contrato", "Segmento", "Número", "Objeto", "Valor", "Prazo", "Status", "Observação"])
        with conn() as c:
            dados = c.execute("""
                SELECT a.*, co.numero contrato, s.nome segmento
                FROM aditivos a
                LEFT JOIN contratos co ON co.id=a.contrato_id
                LEFT JOIN segmentos s ON s.id=a.segmento_id
                ORDER BY a.id DESC
            """).fetchall()
        for r in dados:
            ws.append([r["id"], r["contrato"], r["segmento"], r["numero_aditivo"], r["objeto"], r["valor"], r["prazo"], r["status"], r["observacao"]])
        self._formatar_header(ws, 1)

    def _montar_aba_anexos(self, ws, rows):
        ids = self._ids_filtrados(rows)
        ws.append(["Orçamento", "Tipo", "Arquivo", "Caminho", "Data upload", "Usuário", "Observação"])
        if not ids:
            self._formatar_header(ws, 1); return
        ph = ",".join("?" for _ in ids)
        with conn() as c:
            dados = c.execute(f"""
                SELECT o.codigo, d.*
                FROM documentos d JOIN orcamentos o ON o.id=d.orcamento_id
                WHERE d.orcamento_id IN ({ph})
                ORDER BY o.id DESC, d.id DESC
            """, ids).fetchall()
        for r in dados:
            nome = str(r["nome_arquivo"] or "").replace("_", " ").replace("  ", " ")
            ws.append([r["codigo"], r["tipo"], nome, r["caminho"], r["data_upload"], r["usuario"], r["observacao"]])
        self._formatar_header(ws, 1)

    def _montar_aba_pendencias(self, ws, rows):
        ws.append(["Código", "Status", "Tipo", "Descrição", "Prioridade"])
        pendencias = []
        for r in rows:
            pendencias.extend(self._pendencias_orcamento(r, self.fontes_orcamento(r["id"])))
        for p in pendencias:
            ws.append(p)
        self._formatar_header(ws, 1)
